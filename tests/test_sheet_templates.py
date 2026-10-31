"""New cohort sheets are made from the project's template, and every formula
follows its own animal.

WHY (2026-10-30): the old builder held one layout in code; the lab's real
workbooks had drifted from it, so a new sheet matched none of them, and a
second project could get no correctly formatted sheet at all. A template is a
cohort workbook with what was measured cleared. The builder repeats the model
animal's rows for each new animal, and a formula that pointed at "this
animal's" row (its baseline weight) must point at the NEW animal's row -- a
copied formula that still pointed at the first animal would silently compute
every animal's weight percent from someone else's baseline.
"""
import json
from datetime import datetime

import openpyxl
import pytest

from mousedb import sheet_projects, sheet_templates as st
from mousedb.cohort_tools.make_sheets import NewSheetRefused

A = ["PROJA_05_01", "PROJA_05_02", "PROJA_05_03"]


def _reference(path):
    """A worked-in cohort workbook of three animals."""
    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "Meta"
    meta.append(["SubjectID", "Sex", "Cohort", "Anaesthetic"])
    for a, sex in zip(A, "FMF"):
        meta.append([a, sex, "PROJA_05", "DrugX"])

    w = wb.create_sheet("Weights")
    w.append(["Date", "Day", "Animal", "Weight_g", "Pct"])
    r = 2
    for day in (0, 1, 2):
        for i, a in enumerate(A):
            when = datetime(2026, 7, 1 + day) if day < 2 else None      # day 2: laid out, not done
            w.append([when, day, a, 20 + i if day < 2 else None, "=D%d/$D$%d*100" % (r, 2 + i)])
            r += 1

    t = wb.create_sheet("Trays")
    t.append(["Date", "Animal", "Weight", "Pct", "Phase", "Tray", "S1", "S2", "", "", "Mean"])
    t.cell(1, 11).value = None
    side_row = 2
    r = 2
    for day, phase, tray in ((0, "Flat", "F"), (1, "Pillar", "P")):
        for i, a in enumerate(A):
            for k in (1, 2):
                t.cell(r, 1).value = datetime(2026, 7, 1 + day)
                t.cell(r, 2).value = a
                t.cell(r, 3).value = 20 + i
                t.cell(r, 4).value = "=C%d/'Weights'!$D$%d" % (r, 2 + i)
                t.cell(r, 5).value = phase
                t.cell(r, 6).value = "%s%d" % (tray, k)
                t.cell(r, 7).value = 1
                t.cell(r, 8).value = 2
                r += 1
            # side table (after the empty columns I:J): one row per animal-day
            first = 2 + (day * 3 + i) * 2
            t.cell(side_row, 11).value = a
            t.cell(side_row, 12).value = "=COUNTIF(G%d:H%d,2)" % (first, first + 1)
            side_row += 1
    # cells beside the weights table: a mean over day 0 of every animal
    w.cell(2, 7).value = "Mean"
    w.cell(2, 8).value = "=AVERAGE(D2:D4)"

    plan = wb.create_sheet("Plan")
    plan.append(["Phase", "Start_Day", "Start_Date"])
    plan.append(["Flat", 0, datetime(2026, 7, 1)])
    plan.append(["Pillar", 1, datetime(2026, 7, 2)])
    wb.save(path)


@pytest.fixture
def lab(tmp_path, monkeypatch):
    cfg = tmp_path / "config" / "mousedb_lab_config.json"
    cfg.parent.mkdir()
    cfg.write_text(json.dumps({"sheet_projects": {
        "PROJA": {"sheet_name": "ProjectA_{NN}_Animal_Tracking.xlsx"}}}), encoding="utf-8")
    monkeypatch.setenv(sheet_projects.ENV_VAR, str(cfg))
    ref = tmp_path / "ProjectA_05_Animal_Tracking.xlsx"
    _reference(ref)
    sheets = tmp_path / "sheets"
    sheets.mkdir()
    return ref, sheets


def test_the_template_keeps_the_format_and_clears_what_was_measured(lab):
    ref, _ = lab
    out, rep = st.make_project_template("PROJA", ref, "PROJA_05")
    wb = openpyxl.load_workbook(out)
    t, w, meta, plan = wb["Trays"], wb["Weights"], wb["Meta"], wb["Plan"]
    assert t["E2"].value == "Flat" and t["F3"].value == "F2"          # the schedule stays
    assert t["A2"].value is None and t["C2"].value is None             # dates, weights go
    assert t["G2"].value is None                                        # scores go
    assert t["D2"].value.startswith("=")                               # formulas stay
    assert w["B5"].value == 1                                          # day numbers are schedule
    assert meta["B2"].value is None                                    # sex is a measurement
    assert meta["C2"].value == "PROJA_05"                              # the cohort id is kept
    assert meta["D2"].value is None                                    # same-for-all: cleared ...
    assert "Anaesthetic = 'DrugX'" in rep.text()                       # ... and said
    assert plan["A2"].value == "Flat" and plan["C2"].value is None     # plan kept, its dates not
    assert sheet_projects.definitions()["PROJA"].template == out       # recorded in the settings
    assert (out.parent / "PROJA_template_report.txt").is_file()


def test_a_kept_column_is_kept(lab):
    ref, _ = lab
    out, _ = st.make_project_template("PROJA", ref, "PROJA_05", keep={"Meta": ["Anaesthetic"]})
    assert openpyxl.load_workbook(out)["Meta"]["D2"].value == "DrugX"


def test_every_formula_follows_its_own_animal(lab):
    ref, sheets = lab
    st.make_project_template("PROJA", ref, "PROJA_05")
    target, problems = st.build_sheet("PROJA_07", "3,7", sheets)
    assert target.name == "ProjectA_07_Animal_Tracking.xlsx"
    assert problems == []
    wb = openpyxl.load_workbook(target)
    w, t, meta = wb["Weights"], wb["Trays"], wb["Meta"]
    assert [meta.cell(r, 1).value for r in (2, 3)] == ["PROJA_07_03", "PROJA_07_07"]
    assert meta["C2"].value == "PROJA_07"
    # weights: day-major, each animal's percent against ITS OWN day-0 row
    assert [w.cell(r, 3).value for r in range(2, 8)] == ["PROJA_07_03", "PROJA_07_07"] * 3
    assert w["E2"].value == "=D2/$D$2*100" and w["E3"].value == "=D3/$D$3*100"
    assert w["E5"].value == "=D5/$D$3*100"
    # the mean beside the table covers exactly the new animals' day 0
    assert (w["G2"].value, w["H2"].value) == ("Mean", "=AVERAGE(D2:D3)")
    # trays: two rows per animal per day; weight percent points at its own baseline
    assert [t.cell(r, 2).value for r in range(2, 10)] == (["PROJA_07_03"] * 2 + ["PROJA_07_07"] * 2) * 2
    assert t["D4"].value == "=C4/'Weights'!$D$3"
    assert t["D8"].value == "=C8/'Weights'!$D$3"
    assert (t["E6"].value, t["F7"].value) == ("Pillar", "P2")
    # side table: one row per animal-day, each over that animal's own two rows
    assert [t.cell(r, 11).value for r in range(2, 6)] == ["PROJA_07_03", "PROJA_07_07"] * 2
    assert t["L3"].value == "=COUNTIF(G4:H5,2)" and t["L5"].value == "=COUNTIF(G8:H9,2)"
    assert t["A2"].value is None and t["G2"].value is None
    assert "_template" not in wb.sheetnames


@pytest.mark.parametrize("text, expected", [
    ("01-03", ["PROJA_07_01", "PROJA_07_02", "PROJA_07_03"]),
    ("3, 4,16", ["PROJA_07_03", "PROJA_07_04", "PROJA_07_16"]),
    ("1-2,9", ["PROJA_07_01", "PROJA_07_02", "PROJA_07_09"]),
])
def test_subject_numbers_as_typed(text, expected):
    assert st.parse_animals(text, "PROJA_07") == expected


@pytest.mark.parametrize("bad", ["", "a-b", "3-1", "0", "100", "1,1", "1-3,2"])
def test_subject_numbers_that_cannot_be_meant_are_refused(bad):
    with pytest.raises(NewSheetRefused):
        st.parse_animals(bad, "PROJA_07")


def test_no_template_is_refused_with_the_command_to_make_one(lab):
    _, sheets = lab
    with pytest.raises(NewSheetRefused) as e:
        st.build_sheet("PROJA_07", "01-03", sheets)
    assert "mousedb sheet-template make PROJA" in str(e.value)
    assert list(sheets.iterdir()) == []


def test_an_existing_workbook_or_a_lookalike_is_never_written_over_or_beside(lab):
    ref, sheets = lab
    st.make_project_template("PROJA", ref, "PROJA_05")
    for existing in ("ProjectA_07_Animal_Tracking.xlsx", "Copy of ProjectA_7_Animal_Tracking.xlsx"):
        (sheets / existing).write_bytes(b"real sheet")
        with pytest.raises(NewSheetRefused):
            st.build_sheet("PROJA_07", "01-03", sheets)
        assert (sheets / existing).read_bytes() == b"real sheet"
        assert sorted(p.name for p in sheets.iterdir()) == [existing]
        (sheets / existing).unlink()


@pytest.mark.parametrize("bad", ["PROJB_07", "PROJA_7", "PROJA_07_01", "07", "x PROJA_07"])
def test_a_cohort_name_that_is_not_one_projects_cohort_is_refused(lab, bad):
    ref, sheets = lab
    st.make_project_template("PROJA", ref, "PROJA_05")
    with pytest.raises(NewSheetRefused):
        st.build_sheet(bad, "01-03", sheets)
    assert list(sheets.iterdir()) == []


def test_remaking_a_template_keeps_the_old_one(lab):
    ref, _ = lab
    first, _ = st.make_project_template("PROJA", ref, "PROJA_05")
    st.make_project_template("PROJA", ref, "PROJA_05")
    kept = [p.name for p in first.parent.glob("PROJA_replaced_*.xlsx")]
    assert len(kept) == 1 and first.is_file()


def test_the_command_line_refuses_with_its_own_exit_code(lab, capsys):
    from mousedb.cohort_tools import make_sheets as ms
    _, sheets = lab
    assert ms.main(["--new", "--cohort", "PROJA_07", "--animals", "01-03",
                    "--output-dir", str(sheets)]) == ms.REFUSED_EXIT
    assert "REFUSED:" in capsys.readouterr().err


def test_a_colour_rule_that_maps_onto_itself_is_not_reported(lab):
    """Same number of new animals as the template: a block's colour rule comes
    out identical, which is correct, not a problem."""
    ref, sheets = lab
    wb = openpyxl.load_workbook(ref)
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    wb["Weights"].conditional_formatting.add(
        "D2:D4", CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor="FF0000")))
    wb.save(ref)
    st.make_project_template("PROJA", ref, "PROJA_05")
    target, problems = st.build_sheet("PROJA_07", "01-03", sheets)
    assert problems == []
    assert [str(c.sqref) for c in openpyxl.load_workbook(target)["Weights"].conditional_formatting] == ["D2:D4"]


def test_a_virus_batch_is_emptied_but_its_formulas_stay():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Batch one - 20260826"])
    ws.append(["Virus", "Starting concentration", "Solution Parts Total", "parts for this virus"])
    ws.append(["AAV-A", 100, 6, "=C3/B3"])
    ws.append(["AAV-B", 60, 6, "=C4/B4"])
    ws.append(["Totals", None, None, "=SUM(D3:D4)"])
    assert st._clear_batch_blocks(ws) == 1
    assert ws["A1"].value == "Batch <label> - YYYYMMDD"
    assert ws["A3"].value is None and ws["B3"].value is None and ws["C4"].value is None
    assert ws["D3"].value == "=C3/B3" and ws["D5"].value == "=SUM(D3:D4)"
    assert ws["A5"].value == "Totals"
