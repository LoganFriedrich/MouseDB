"""What needs a person is gathered and shown, and written beside each sheet.

WHY (Logan, 2026-10-26): "There needs to be a way to alert a human to an
issue of the kind you are hunting down." -- a session with no hand scores,
a date the import read as a typo, a failed import. Shown in red on Where Is
My Data, and written per cohort into "!! <cohort> - FIX THESE.xlsx" beside its
sheet, where the sheet is fixed.
"""
import json
import os

import openpyxl

# The suite's other Qt test does the same: no screen is needed, or available.
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from mousedb import attention, sheet_sync


def _ledger(tmp_path, monkeypatch, entries):
    p = tmp_path / "ledger.jsonl"
    p.write_text("\n".join(json.dumps(e) for e in entries) + "\n", encoding="utf-8")
    monkeypatch.setattr(sheet_sync, "LEDGER", p)


def test_the_pipelines_lists_and_the_last_import_are_gathered(tmp_path, monkeypatch):
    att = tmp_path / "pipe" / "attention"
    att.mkdir(parents=True)
    (att / "HOST_hand_scores.json").write_text(json.dumps({"items": [
        {"needs_person": True, "cohort_id": "CNT_02", "says": "20250909_CNT0216_P1 has none"},
        {"needs_person": False, "cohort_id": "CNT_02", "says": "only waiting"}]}))
    _ledger(tmp_path, monkeypatch, [
        {"cohort_id": "CNT_05", "success": True, "dry_run": False,
         "warnings": ["REPAIRED CNT_05 3b_Manual_Tray: date 2025-09-28 read as 2026-09-28 -- x",
                      "an ordinary warning"]},
        {"cohort_id": "CNT_06", "success": False, "dry_run": False, "error": "sheet locked"},
        {"cohort_id": "CNT_06", "success": False, "dry_run": True, "error": "a dry run"},
    ])
    mr = attention.mousereach_items(tmp_path / "pipe")
    assert [i["says"] for i in mr] == ["20250909_CNT0216_P1 has none"]
    sh = attention.sheet_items(["CNT_05", "CNT_06"])
    assert [i["cohort_id"] for i in sh] == ["CNT_05", "CNT_06"]
    assert sh[0]["says"].startswith("REPAIRED CNT_05")
    assert "failed: sheet locked" in sh[1]["says"]


def test_the_fix_these_file_is_written_and_removed_when_nothing_is_left(tmp_path):
    items = [{"cohort_id": "CNT_05", "source": "sheet import", "says": "REPAIRED ..."}]
    p = attention.write_fix_these("CNT_05", items, tmp_path)
    assert p.name == "!! CNT_05 - FIX THESE.xlsx"
    ws = openpyxl.load_workbook(p).active
    assert ws["A1"].value == "CNT_05: 1 thing(s) to fix"
    assert ws["A5"].value == "REPAIRED ..."
    # never mistaken for the cohort's sheet
    from mousedb.cohort_sheets import _SHEET_RE
    assert not _SHEET_RE.match(p.name)
    assert attention.write_fix_these("CNT_05", [], tmp_path) is None
    assert not p.exists()


def test_the_tab_shows_them_in_red_and_hides_when_none():
    from PyQt5.QtWidgets import QApplication, QLabel
    from types import SimpleNamespace
    app = QApplication.instance() or QApplication([])   # kept: a dropped app is collected
    from mousedb.gui.data_status_tab import DataStatusTab
    tab = SimpleNamespace(attention_label=QLabel(), ATTENTION_SHOWN=DataStatusTab.ATTENTION_SHOWN)
    DataStatusTab._show_attention(tab, [{"says": "fix the sheet"}])
    assert not tab.attention_label.isHidden()
    assert "Needs your attention (1)" in tab.attention_label.text()
    DataStatusTab._show_attention(tab, [])
    assert tab.attention_label.isHidden()
    del app


def test_a_filed_video_whose_session_has_no_hand_scores_is_listed(tmp_path):
    import pandas as pd
    pd.DataFrame({
        "video_name": ["20250909_CNT0216_P1", "20250909_CNT0215_P1", "20220811_ASPA1011_P1",
                       "20250909_CNT0216_E1"],
        "subject_id": ["CNT_02_16", "CNT_02_15", "ASPA_10_11", "CNT_02_16"],
        "session_date": ["2025-09-09"] * 3 + ["2025-09-09"],
        "tray_type": ["P", "P", "P", "E"], "run_number": [1, 1, 1, 1],
    }).to_parquet(tmp_path / "video_pellets.parquet")
    hand = tmp_path / "hand.csv"
    hand.write_text("video_key,pellet,score\n20250909_CNT0215_P1,1,2\n", encoding="utf-8")
    found = attention.filed_without_scores(snapshot_dir=tmp_path, hand_csv=hand)
    # the scored one, the frozen-sheet project and the Easy tray are not listed
    assert [i["cohort_id"] for i in found] == ["CNT_02"]
    assert found[0]["says"].startswith("20250909_CNT0216_P1 is in the database but its "
                                       "session has no hand scores")
