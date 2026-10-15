"""Why this file exists: "New cohort sheet..." must never cost the lab a sheet.

That button (gui.sheets_tab.new_sheet -> cohort_tools.make_sheets) is the one
thing mousedb is BUILT to write into the folder of source-of-record tracking
workbooks, and it could replace a live, filled sheet with an empty one: the
cohort number in the file name was whatever followed the LAST underscore of the
typed text (so another project's id, an animal id, or any typo ending "_01"
named cohort 01's existing workbook), nothing checked whether that file
existed, and a blocked save force-closed every Excel window on the machine.

These tests pin the fix at both layers:
  * make_sheets refuses, whoever calls it -- a name not exactly PREFIX_NN, a
    cohort that already has a workbook (exact name or a variant), a file that
    appears while the sheet is being built -- and writes nothing;
  * the Tracking Sheets tab refuses up front in a message box, shows the
    builder's own refusals, and never lets an exception take the window down;
  * and the tab's other fixes from the same review: a pin that fails is a
    message, not a crash, and broken pins and files refused as sheets are shown
    in plain view, not only in a terminal.

Every folder is a pytest tmp_path and the config file is redirected; nothing
real is read or written, and no real subprocess is started.
"""
import inspect
import json
import os
import subprocess
import time
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from mousedb import cohort_sheets as cs  # noqa: E402
from mousedb import sheet_sync as ss  # noqa: E402
from mousedb.cohort_tools import make_sheets as ms  # noqa: E402

REAL = b"the lab's real, filled-in sheet"


@pytest.fixture(autouse=True)
def _hermetic(tmp_path, monkeypatch):
    """No real study facts, config, sheets folder or ledger; the builder's
    'other cohorts in the working directory' scan sees an empty folder."""
    monkeypatch.setenv("MOUSEDB_STUDY_FACTS", str(tmp_path / "facts.json"))
    cwd = tmp_path / "cwd"
    cwd.mkdir()
    monkeypatch.setattr(ms, "SCRIPT_DIR", cwd)
    monkeypatch.setattr(cs, "CONFIG_PATH", tmp_path / "config.json")
    monkeypatch.delenv(cs.ASPA_ENV_VAR, raising=False)
    monkeypatch.setattr(ss, "LEDGER", tmp_path / "ledger.jsonl")


@pytest.fixture
def sheets(tmp_path, monkeypatch):
    """A sheets folder holding cohort 01's real sheet, configured as THE folder."""
    d = tmp_path / "sheets"
    d.mkdir()
    (d / "Connectome_01_Animal_Tracking.xlsx").write_bytes(REAL)
    monkeypatch.setenv(cs.ENV_VAR, str(d))
    return d


def _snapshot(folder: Path) -> dict:
    return {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in folder.iterdir()}


# ---------------------------------------------------------------------------
# the name is read strictly
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("text, cohort_id, number, project_type, filename", [
    ("CNT_07", "CNT_07", "07", "cnt", "Connectome_07_Animal_Tracking.xlsx"),
    ("  cnt_07 ", "CNT_07", "07", "cnt", "Connectome_07_Animal_Tracking.xlsx"),
    ("ENCR_01", "ENCR_01", "01", "encr", "Enhancer_01_Animal_Tracking.xlsx"),
    ("CNT_00", "CNT_00", "00", "cnt", "Connectome_00_Animal_Tracking.xlsx"),
])
def test_a_cohort_name_is_prefix_underscore_two_digits(text, cohort_id, number,
                                                       project_type, filename):
    c = ms.parse_cohort_name(text)
    assert (c.cohort_id, c.number, c.project_type, c.filename) == (
        cohort_id, number, project_type, filename)


@pytest.mark.parametrize("text", [
    "CNTS_01",        # another project's id ending in _01: the original hazard
    "OTHER_01",
    "CNT_05_01",      # an animal id
    "01", "1",        # no project
    "CNT_1", "CNT_001", "CNT07", "CNT-07", "CNT__07", "CNT_07.xlsx",
    "Connectome_01",  # the file word, not a project prefix
    "CNT_\u0660\u0667",  # digits from another script
    "", "   ", None,
])
def test_anything_else_is_refused_saying_what_is_expected(text):
    with pytest.raises(ms.NewSheetRefused) as e:
        ms.parse_cohort_name(text)
    assert "CNT_07" in str(e.value)   # the message shows the expected form


# ---------------------------------------------------------------------------
# the builder never writes over, or beside, an existing workbook
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("typed", ["CNT_01", "cnt_01", "CNTS_01", "OTHER_01",
                                   "CNT_05_01", "01"])
def test_cohort_01s_sheet_survives_every_name_that_could_reach_it(sheets, typed):
    before = _snapshot(sheets)
    with pytest.raises(ms.NewSheetRefused):
        ms.create_new_cohort_file(typed, "2026-01-05", 1, sheets)
    assert _snapshot(sheets) == before   # same files, same bytes, same times


@pytest.mark.parametrize("existing", [
    "Connectome_07_Animal_Tracking1.xlsx",          # a numbered variant
    "Connectome_07_Animal_Tracking (2).xlsx",       # a sync / download copy
    "Connectome_7_Animal_Tracking.xlsx",            # number without its zero
    "connectome_07_animal_tracking.xlsx",           # other letter case
])
def test_a_variant_of_the_cohorts_name_also_blocks_a_new_sheet(tmp_path, existing):
    """An empty workbook beside a variant would be the NEWEST candidate for
    the cohort -- the file an unpinned import reads."""
    d = tmp_path / "folder"
    d.mkdir()
    (d / existing).write_bytes(REAL)
    before = _snapshot(d)
    with pytest.raises(ms.NewSheetRefused) as e:
        ms.create_new_cohort_file("CNT_07", "2026-01-05", 1, d)
    assert existing in str(e.value)
    assert _snapshot(d) == before


def test_other_cohorts_files_do_not_block(tmp_path):
    d = tmp_path / "folder"
    d.mkdir()
    for name in ("Connectome_17_Animal_Tracking.xlsx", "Connectome_70_Animal_Tracking.xlsx",
                 "Copy of Connectome_07_Animal_Tracking.xlsx"):
        (d / name).write_bytes(REAL)
    p = ms.create_new_cohort_file("CNT_07", "2026-01-05", 1, d)
    assert p == d / "Connectome_07_Animal_Tracking.xlsx" and p.is_file()


def test_a_new_sheet_is_one_file_with_the_validated_ids(tmp_path):
    import openpyxl
    d = tmp_path / "folder"
    d.mkdir()
    p = ms.create_new_cohort_file(" cnt_09 ", "2026-01-05", 2, d)
    # exactly one file: no temporary name left behind
    assert sorted(x.name for x in d.iterdir()) == ["Connectome_09_Animal_Tracking.xlsx"]
    ws = openpyxl.load_workbook(p)["0a_Metadata"]
    assert [ws["A2"].value, ws["A3"].value] == ["CNT_09_01", "CNT_09_02"]


def test_a_workbook_that_appears_during_the_save_is_left_alone(tmp_path, monkeypatch):
    """Someone (a sync client, a colleague) puts the cohort's workbook in the
    folder while the new one is being written: theirs stays, nothing of ours
    is left behind."""
    d = tmp_path / "folder"
    d.mkdir()
    target = d / "Connectome_07_Animal_Tracking.xlsx"
    real_save = ms.openpyxl.Workbook.save

    def save_while_someone_else_writes(self, filename):
        target.write_bytes(REAL)
        return real_save(self, filename)

    monkeypatch.setattr(ms.openpyxl.Workbook, "save", save_while_someone_else_writes)
    with pytest.raises(ms.NewSheetRefused) as e:
        ms.create_new_cohort_file("CNT_07", "2026-01-05", 1, d)
    assert "appeared" in str(e.value)
    assert target.read_bytes() == REAL
    assert [x.name for x in d.iterdir()] == [target.name]


def test_a_workbook_that_appears_during_the_build_is_left_alone(tmp_path, monkeypatch):
    """The build takes a while; the folder is checked again before saving, and
    a variant name counts there too."""
    d = tmp_path / "folder"
    d.mkdir()
    intruder = d / "Connectome_07_Animal_Tracking1.xlsx"
    real_autoscale = ms.autoscale_columns

    def autoscale_while_someone_else_writes(ws):
        intruder.write_bytes(REAL)
        return real_autoscale(ws)

    monkeypatch.setattr(ms, "autoscale_columns", autoscale_while_someone_else_writes)
    with pytest.raises(ms.NewSheetRefused):
        ms.create_new_cohort_file("CNT_07", "2026-01-05", 1, d)
    assert [x.name for x in d.iterdir()] == [intruder.name]
    assert intruder.read_bytes() == REAL


@pytest.mark.parametrize("kwargs", [
    {"project_type": "encr"},              # disagrees with the CNT prefix
    {"start_date": "2026-13-40"},
    {"start_date": "next monday"},
    {"num_mice": 0}, {"num_mice": 100}, {"num_mice": "many"},
])
def test_other_bad_requests_are_refused_before_anything_is_written(tmp_path, kwargs):
    d = tmp_path / "folder"
    d.mkdir()
    args = dict(cohort_name="CNT_08", start_date="2026-01-05", num_mice=1, output_dir=d)
    args.update(kwargs)
    with pytest.raises(ms.NewSheetRefused):
        ms.create_new_cohort_file(**args)
    assert list(d.iterdir()) == []


def test_the_builder_never_closes_excel_or_waits_for_a_keypress():
    """The old save loop ran 'taskkill /F /IM EXCEL.EXE' (every Excel window on
    the machine, unsaved work lost) and then waited on input(), which a GUI's
    subprocess can never answer."""
    src = inspect.getsource(ms)
    assert "taskkill" not in src.lower()
    assert "EXCEL.EXE" not in src.upper()
    assert "input(" not in inspect.getsource(ms.create_new_cohort_file)


def test_the_command_line_refusal_is_one_line_and_its_own_exit_code(sheets, capsys):
    before = _snapshot(sheets)
    for cohort in ("CNTS_01", "CNT_01"):
        rc = ms.main(["--new", "--cohort", cohort, "--start-date", "2026-01-05",
                      "--output-dir", str(sheets)])
        assert rc == ms.REFUSED_EXIT
        err = capsys.readouterr().err
        assert err.startswith("REFUSED: ") and "Nothing was written." in err
        assert "Traceback" not in err
    assert _snapshot(sheets) == before


def test_fix_mode_never_writes_over_its_own_input(tmp_path):
    """--fix keeps the input's name, so an --output-dir naming the input's own
    folder used to replace the sheet in place. The input is a REAL workbook:
    the old code fails to load a few bytes of text and so never reaches the
    save, which would let this test pass against the very code it guards."""
    import openpyxl
    d = tmp_path / "live"
    d.mkdir()
    sheet = d / "Connectome_01_Animal_Tracking.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "0a_Metadata"
    wb.active["A1"], wb.active["A2"] = "SubjectID", "CNT_01_01"
    wb.save(sheet)
    before = _snapshot(d)
    assert ms.fix_existing_file(sheet, output_dir=d) is None
    assert _snapshot(d) == before


@pytest.mark.skipif(os.name == "nt", reason="file modes are POSIX; Windows has ACLs")
def test_a_new_sheet_is_readable_by_colleagues_on_a_shared_posix_folder(tmp_path):
    """The workbook is written under a temporary name and renamed; the
    temporary file must not be created owner-only (mkstemp's 0600), or the
    published sheet would be closed to everyone else in the lab."""
    d = tmp_path / "folder"
    d.mkdir()
    um = os.umask(0o022)
    try:
        p = ms.create_new_cohort_file("CNT_07", "2026-01-05", 1, d)
    finally:
        os.umask(um)
    assert p.stat().st_mode & 0o777 == 0o644


# ---------------------------------------------------------------------------
# the Tracking Sheets tab
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def qapp():
    from PyQt5.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _settle(tab, timeout=20.0):
    """Let the tab's background command(s) finish and their results land."""
    from PyQt5.QtWidgets import QApplication
    app = QApplication.instance()
    end = time.time() + timeout
    while time.time() < end:
        app.processEvents()
        if not tab._busy and all(t.isFinished() for t in tab._threads):
            app.processEvents()
            if not tab._busy:
                return
        time.sleep(0.01)
    raise AssertionError("the tab's background command did not finish")


@pytest.fixture
def ui(qapp, monkeypatch):
    """Records every message box and answers the dialogs from queues, so a
    test can drive the tab like a person and see what the person would see."""
    from PyQt5.QtWidgets import QFileDialog, QInputDialog, QMessageBox

    class UI:
        boxes = []          # (kind, title, text)
        texts = []          # answers for QInputDialog.getText, in order
        items = []          # answers for QInputDialog.getItem, in order
        asked = []          # prompts shown
        answer = QMessageBox.Yes
        folder = ""

    def box(kind):
        def show(parent, title, text, *a, **k):
            UI.boxes.append((kind, title, text))
            return QMessageBox.Ok
        return staticmethod(show)

    for kind in ("warning", "information", "critical"):
        monkeypatch.setattr(QMessageBox, kind, box(kind))

    def question(parent, title, text, *a, **k):
        UI.boxes.append(("question", title, text))
        return UI.answer

    def get_text(parent, title, label, *a, **k):
        UI.asked.append(label)
        return (UI.texts.pop(0), True) if UI.texts else ("", False)

    def get_item(parent, title, label, items, *a, **k):
        UI.asked.append(label)
        UI.offered = list(items)
        return (UI.items.pop(0), True) if UI.items else ("", False)

    monkeypatch.setattr(QMessageBox, "question", staticmethod(question))
    monkeypatch.setattr(QInputDialog, "getText", staticmethod(get_text))
    monkeypatch.setattr(QInputDialog, "getItem", staticmethod(get_item))
    monkeypatch.setattr(QFileDialog, "getExistingDirectory",
                        staticmethod(lambda *a, **k: UI.folder))
    UI.boxes, UI.texts, UI.items, UI.asked = [], [], [], []
    return UI


@pytest.fixture
def make_tab(qapp, monkeypatch):
    """A tab whose status comes from ``status`` (a dict, or a function that
    returns one) instead of the machine's real folder."""
    from mousedb.gui import sheets_tab as st_mod
    made = []

    def make(status):
        calls = []

        def fake_status():
            calls.append(1)
            return status() if callable(status) else json.loads(json.dumps(status))

        monkeypatch.setattr(st_mod, "_status_with_broken_pins", fake_status)
        tab = st_mod.TrackingSheetsTab()
        _settle(tab)
        tab.status_calls = calls
        made.append(tab)
        return tab

    yield make
    for tab in made:
        _settle(tab)
        tab.deleteLater()
    qapp.processEvents()


def _click(tab, text):
    from PyQt5.QtWidgets import QPushButton
    buttons = [b for b in tab.findChildren(QPushButton) if b.text() == text]
    assert len(buttons) == 1, text
    buttons[0].click()
    _settle(tab)


def _no_subprocess(monkeypatch):
    def refuse(*a, **k):
        raise AssertionError("the sheet builder must not be started")
    monkeypatch.setattr(subprocess, "run", refuse)


def _folder_status(d, cohorts=(), **extra):
    st = {"cnt_sheets_dir": str(d), "ledger": "ledger", "cohorts": list(cohorts),
          "refused_files": [], "problem": None}
    st.update(extra)
    return st


@pytest.mark.parametrize("typed, says", [
    ("CNTS_01", "does not start with a project"),
    ("CNT_05_01", "is not a cohort name"),
    ("01", "is not a cohort name"),
    ("CNT_01", "already has a workbook"),
    ("ENCR_02", "not a CNT cohort"),       # this folder is read for CNT sheets only
])
def test_the_tab_refuses_in_a_message_box_and_starts_nothing(sheets, make_tab, ui,
                                                              monkeypatch, typed, says):
    tab = make_tab(_folder_status(sheets))
    _no_subprocess(monkeypatch)
    before = _snapshot(sheets)
    ui.texts = [typed, "2026-01-05"]
    _click(tab, "New cohort sheet...")
    kinds = [b[0] for b in ui.boxes]
    assert kinds == ["warning"], ui.boxes
    assert says in ui.boxes[0][2] and "Nothing was written." in ui.boxes[0][2]
    assert len(ui.asked) == 1            # never got as far as asking for a date
    assert _snapshot(sheets) == before


def test_the_tab_refuses_a_bad_date_before_starting_anything(sheets, make_tab, ui,
                                                             monkeypatch):
    tab = make_tab(_folder_status(sheets))
    _no_subprocess(monkeypatch)
    ui.texts = ["CNT_07", "31/01/2026"]
    _click(tab, "New cohort sheet...")
    assert [b[0] for b in ui.boxes] == ["warning"]
    assert "is not a date" in ui.boxes[0][2]


def test_the_builders_own_refusal_reaches_the_person(sheets, make_tab, ui, monkeypatch):
    """Something can change between the tab's check and the build; the
    builder's REFUSED line is shown, not swallowed or shown as a crash."""
    tab = make_tab(_folder_status(sheets))
    seen = {}

    def fake_run(cmd, **kw):
        seen["cmd"] = cmd
        return subprocess.CompletedProcess(cmd, ms.REFUSED_EXIT, stdout="",
                                           stderr="REFUSED: CNT_07 already has a workbook "
                                                  "in the folder: X.\nNothing was written.\n")

    monkeypatch.setattr(subprocess, "run", fake_run)
    ui.texts = ["CNT_07", "2026-01-05"]
    _click(tab, "New cohort sheet...")
    cmd = seen["cmd"]
    assert cmd[cmd.index("--cohort") + 1] == "CNT_07"
    assert cmd[cmd.index("--output-dir") + 1] == str(sheets)
    assert [b[0] for b in ui.boxes] == ["question", "warning"]
    assert "CNT_07 already has a workbook" in ui.boxes[1][2]
    assert tab.status_calls and len(tab.status_calls) >= 2     # table re-read


def test_the_tab_says_so_when_the_sheet_is_created(sheets, make_tab, ui, monkeypatch):
    tab = make_tab(_folder_status(sheets))

    def fake_run(cmd, **kw):
        (sheets / "Connectome_07_Animal_Tracking.xlsx").write_bytes(b"new")
        return subprocess.CompletedProcess(cmd, 0, stdout="[OK] created\n", stderr="")

    monkeypatch.setattr(subprocess, "run", fake_run)
    ui.texts = ["CNT_07", "2026-01-05"]
    _click(tab, "New cohort sheet...")
    assert [b[0] for b in ui.boxes] == ["question", "information"]
    assert "Created Connectome_07_Animal_Tracking.xlsx" in ui.boxes[1][2]
    assert (sheets / "Connectome_01_Animal_Tracking.xlsx").read_bytes() == REAL


def test_a_sheets_folder_gone_since_the_refresh_is_never_recreated(tmp_path, make_tab, ui,
                                                                   monkeypatch):
    """The builder creates a missing output folder. A sheets folder moved or
    offline since the last Refresh must be refused -- not recreated at its old
    path holding one empty sheet."""
    gone = tmp_path / "moved_away"
    tab = make_tab(_folder_status(gone))
    _no_subprocess(monkeypatch)
    ui.texts = ["CNT_07", "2026-01-05"]
    _click(tab, "New cohort sheet...")
    assert [b[0] for b in ui.boxes] == ["warning"]
    assert "not there any more" in ui.boxes[0][2] and "Nothing was written." in ui.boxes[0][2]
    assert not gone.exists()


def test_a_sheets_folder_gone_while_the_confirmation_was_open(tmp_path, make_tab, ui,
                                                             monkeypatch):
    from PyQt5.QtWidgets import QMessageBox
    d = tmp_path / "sheets_here"
    d.mkdir()
    tab = make_tab(_folder_status(d))
    _no_subprocess(monkeypatch)

    def answer_after_the_folder_went(parent, title, text, *a, **k):
        ui.boxes.append(("question", title, text))
        d.rmdir()
        return QMessageBox.Yes

    monkeypatch.setattr(QMessageBox, "question", staticmethod(answer_after_the_folder_went))
    ui.texts = ["CNT_07", "2026-01-05"]
    _click(tab, "New cohort sheet...")
    assert [b[0] for b in ui.boxes] == ["question", "warning"]
    assert "not there any more" in ui.boxes[1][2]
    assert not d.exists()


def test_a_command_that_calls_sys_exit_still_hands_back_a_result(qapp):
    """A worker that let SystemExit out would send no result (the tab stays
    'busy' for good) and PyQt would abort the application."""
    from mousedb.gui.sheets_tab import _Worker
    got = []
    w = _Worker(7, lambda: __import__("sys").exit(2))
    w.done.connect(lambda job, r: got.append((job, r)))
    w.run()          # on this thread: the signal is delivered at once
    assert got and got[0][0] == 7 and "SystemExit" in got[0][1]["problem"]


def test_saying_no_at_the_confirmation_starts_nothing(sheets, make_tab, ui, monkeypatch):
    from PyQt5.QtWidgets import QMessageBox
    tab = make_tab(_folder_status(sheets))
    _no_subprocess(monkeypatch)
    ui.answer = QMessageBox.No
    ui.texts = ["CNT_07", "2026-01-05"]
    _click(tab, "New cohort sheet...")
    assert [b[0] for b in ui.boxes] == ["question"]
    assert "Connectome_07_Animal_Tracking.xlsx" in ui.boxes[0][2]   # the exact file named


def test_a_pin_that_cannot_be_made_is_a_message_not_a_crash(sheets, make_tab, ui):
    """The file chosen was renamed since the table was read: pin_cohort_sheet
    raises ValueError, which used to escape the slot and abort the GUI."""
    (sheets / "Connectome_05_Animal_Tracking.xlsx").write_bytes(REAL)
    row = {"cohort_id": "CNT_05", "project": "CNT", "state": "up_to_date",
           "sheet": "Connectome_05_Animal_Tracking.xlsx", "ambiguous": True,
           "candidates": [{"name": "Connectome_05_Animal_Tracking1.xlsx", "edited": "-"},
                          {"name": "Connectome_05_Animal_Tracking.xlsx", "edited": "-"}]}
    tab = make_tab(_folder_status(sheets, [row]))
    tab.table.selectRow(0)
    ui.items = ["Connectome_05_Animal_Tracking1.xlsx"]     # gone from the folder now
    _click(tab, "This is the sheet")
    assert [b[0] for b in ui.boxes] == ["warning"]
    assert "Nothing was pinned" in ui.boxes[0][2]
    assert cs.pinned_sheet("CNT_05") is None
    assert len(tab.status_calls) >= 2                      # the table was re-read


def test_an_unexpected_error_in_a_button_does_not_escape(sheets, make_tab, ui, monkeypatch):
    tab = make_tab(_folder_status(sheets))

    def broken(path):
        raise RuntimeError("something nobody expected")

    monkeypatch.setattr(os, "startfile", broken, raising=False)
    _click(tab, "Open folder")          # would abort the application if it escaped
    assert [b[0] for b in ui.boxes] == ["warning"]
    assert "something nobody expected" in ui.boxes[0][2]
    assert tab._guarded(lambda: 1 / 0) is None


def test_broken_pins_and_refused_files_are_shown_in_plain_view(tmp_path, make_tab, ui):
    st = _folder_status(
        tmp_path,
        [{"cohort_id": "CNT_05", "project": "CNT", "state": "up_to_date",
          "sheet": "Connectome_05_Animal_Tracking2.xlsx", "candidates": [],
          "broken_pin": "CNT_05 is pinned to A.xlsx, which is gone; read B instead."}],
        all_broken_pins=[
            {"cohort_id": "CNT_05", "pinned": "A.xlsx", "reading": "B",
             "why": "CNT_05 is pinned to A.xlsx, which is gone; read B instead."},
            {"cohort_id": "CNT_06", "pinned": "C.xlsx", "reading": None,
             "why": "CNT_06 is pinned to C.xlsx, which is gone; it has no file."}],
        refused_files=[{"file": "Copy of Connectome_05_Animal_Tracking.xlsx",
                        "reason": "a copy"}])
    tab = make_tab(st)
    text = tab.notices.text()
    assert not tab.notices.isHidden()
    assert text.count("CNT_05 is pinned to A.xlsx") == 1   # said once, not twice
    assert "CNT_06 is pinned to C.xlsx" in text            # the cohort with no row
    assert "Copy of Connectome_05_Animal_Tracking.xlsx -- a copy" in text
    assert "PIN BROKEN" in tab.table.item(0, 2).text()
    assert "[pinned]" not in tab.table.item(0, 2).text()


def test_the_red_box_is_hidden_when_there_is_nothing_to_say(tmp_path, make_tab, ui):
    tab = make_tab(_folder_status(tmp_path, [{"cohort_id": "CNT_05", "project": "CNT",
                                              "state": "up_to_date", "candidates": []}]))
    assert tab.notices.isHidden()


def test_a_real_broken_pin_reaches_the_red_box(sheets, make_tab, ui):
    """End to end through the real status code: a pinned file renamed away
    (cohort still has a row), a pinned cohort left with no file at all (no
    row), and a copy refused as a sheet -- all three said in the tab."""
    from mousedb.gui import sheets_tab as st_mod
    for name in ("Connectome_05_Animal_Tracking.xlsx", "Connectome_05_Animal_Tracking1.xlsx",
                 "Connectome_06_Animal_Tracking.xlsx"):
        (sheets / name).write_bytes(REAL)
    cs.pin_cohort_sheet("CNT_05", "Connectome_05_Animal_Tracking1.xlsx")
    cs.pin_cohort_sheet("CNT_06", "Connectome_06_Animal_Tracking.xlsx")
    (sheets / "Connectome_05_Animal_Tracking1.xlsx").unlink()
    (sheets / "Connectome_06_Animal_Tracking.xlsx").unlink()
    (sheets / "Copy of Connectome_05_Animal_Tracking.xlsx").write_bytes(REAL)

    real = st_mod._status_with_broken_pins
    tab = make_tab(real)
    text = tab.notices.text()
    assert "Connectome_05_Animal_Tracking1.xlsx" in text
    assert "CNT_06" in text
    assert "Copy of Connectome_05_Animal_Tracking.xlsx" in text
    assert text.count("PIN BROKEN") == 2


def test_a_set_folder_holding_only_refused_files_still_shows_them(tmp_path, make_tab, ui,
                                                                  monkeypatch):
    """status() accepts a folder only while it holds a tracking sheet and then
    lists nothing in it. When every sheet there was renamed to a refused name,
    the tab still says which files are there and which pin they broke."""
    from mousedb.gui import sheets_tab as st_mod
    d = tmp_path / "sheets"
    d.mkdir()
    (d / "Connectome_05_Animal_Tracking.xlsx").write_bytes(REAL)
    monkeypatch.setenv(cs.ENV_VAR, str(d))
    cs.pin_cohort_sheet("CNT_05", "Connectome_05_Animal_Tracking.xlsx")
    (d / "Connectome_05_Animal_Tracking.xlsx").rename(
        d / "Copy of Connectome_05_Animal_Tracking.xlsx")
    tab = make_tab(st_mod._status_with_broken_pins)
    text = tab.notices.text()
    assert not tab.notices.isHidden()
    assert "NOT READ AS A SHEET: Copy of Connectome_05_Animal_Tracking.xlsx" in text
    assert "PIN BROKEN: CNT_05" in text
    assert str(d) in tab.folder_label.text()


def test_forget_a_pin_from_the_tab(sheets, make_tab, ui):
    (sheets / "Connectome_06_Animal_Tracking.xlsx").write_bytes(REAL)
    cs.pin_cohort_sheet("CNT_06", "Connectome_06_Animal_Tracking.xlsx")
    (sheets / "Connectome_06_Animal_Tracking.xlsx").unlink()
    from mousedb.gui import sheets_tab as st_mod
    tab = make_tab(st_mod._status_with_broken_pins)
    assert not tab.notices.isHidden()
    offer = "CNT_06  (pinned to Connectome_06_Animal_Tracking.xlsx -- BROKEN)"
    ui.items = [offer]
    _click(tab, "Forget a pin...")
    assert ui.offered == [offer]        # the cohort with no row is still offered
    assert cs.pinned_sheet("CNT_06") is None
    assert tab.notices.isHidden()


def test_set_folder_refuses_a_folder_that_holds_no_sheet(tmp_path, make_tab, ui):
    d = tmp_path / "not_it"
    d.mkdir()
    (d / "Copy of Connectome_05_Animal_Tracking.xlsx").write_bytes(REAL)
    tab = make_tab(_folder_status(tmp_path))
    ui.folder = str(d)
    _click(tab, "Set sheets folder...")
    assert [b[0] for b in ui.boxes] == ["warning"]
    assert "was not changed" in ui.boxes[0][2]
    assert "Copy of Connectome_05_Animal_Tracking.xlsx" in ui.boxes[0][2]
    assert not cs.CONFIG_PATH.exists() or "cohort_sheets_dir" not in json.loads(
        cs.CONFIG_PATH.read_text())


def test_set_folder_says_when_an_environment_variable_overrides_it(sheets, tmp_path, make_tab,
                                                                   ui):
    """The environment variable (here: the 'sheets' fixture) comes before the
    recorded setting, so choosing another folder would change nothing -- and
    used to say only "recorded"."""
    other = tmp_path / "other_sheets"
    other.mkdir()
    (other / "Connectome_02_Animal_Tracking.xlsx").write_bytes(REAL)
    tab = make_tab(_folder_status(sheets))
    ui.folder = str(other)
    _click(tab, "Set sheets folder...")
    assert [b[0] for b in ui.boxes] == ["warning"]
    assert cs.ENV_VAR in ui.boxes[0][2] and str(sheets) in ui.boxes[0][2]
    assert json.loads(cs.CONFIG_PATH.read_text())["cohort_sheets_dir"] == str(other)
