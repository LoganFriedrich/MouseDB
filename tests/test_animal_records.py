"""animal_records copies a tracking sheet's animal-level tabs column for column.

WHY: the per-reach ODC export repeats each animal's surgery details on every row;
the typed surgery table kept only a few of those columns, so everything else the
lab typed into the sheet never reached an export.
"""
from datetime import datetime

import openpyxl
import pytest

from mousedb import animal_records as ar


def _cnt_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4_Contusion_Injury_Details"
    ws.append(["Subject_ID", "Surgery_Date", "Surgery_Type", "Anesthetic", "Intended_kd", "Actual_kd", "Survived"])
    ws.append(["PROJA_01_01", datetime(2025, 7, 2), "Contusion", "drug-x", 60, 65.0, "Y"])
    ws.append(["PROJA_01_02", datetime(2025, 7, 2), "Contusion", "drug-x", 60, None, None])  # planned, not done
    ws.append(["PROJB_01_01", datetime(2025, 7, 2), "Contusion", "drug-x", 60, 61, "Y"])     # other cohort
    sc = wb.create_sheet("5_SC_Injection_Details")
    sc.append(["Subject_ID", "Surgery_Date", "Injected_Virus", "Survived"])
    sc.append(["PROJA_01_01", datetime(2025, 8, 1), "virus-a", "Y"])
    sc.append(["PROJA_01_01", datetime(2025, 9, 1), "virus-b", "N"])
    meta = wb.create_sheet("0a_Metadata")
    meta.append(["SubjectID", "Sex", "Notes"])
    meta.append(["PROJA_01_01", "F", "?"])
    tray = wb.create_sheet("3b_Manual_Tray")
    tray.append(["Date", "Animal", "Sex", "1"])
    tray.append([datetime(2025, 6, 19), "PROJA_01_03", None, 2])
    tray.append([datetime(2025, 6, 20), "PROJA_01_03", "Female", 1])
    tray.append([datetime(2025, 6, 21), "PROJA_01_03", "Male", 1])   # first value wins
    wb.save(path)


def _odc_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ODC"
    ws.append([1, 1, 1, 1])
    ws.append([None, 123, "a note"])
    ws.append(["Group ID", "Animal ID", "Surgery_Date", "Organism_Strain"])
    ws.append(["K", 1, datetime(2022, 9, 22), "strain-k"])
    ws.append(["K", "?", None, "strain-k"])
    wb.save(path)


def test_cnt_tabs_copied_as_written(tmp_path):
    p = tmp_path / "PROJA_01.xlsx"
    _cnt_workbook(p)
    w = []
    recs = ar.read_cnt_workbook(p, "PROJA_01", w)
    got = {(r["subject_id"], r["source_tab"], r["record_no"], r["field"]): r["value"] for r in recs}
    assert got[("PROJA_01_01", "4_Contusion_Injury_Details", 1, "Surgery_Date")] == "2025-07-02"
    assert got[("PROJA_01_01", "4_Contusion_Injury_Details", 1, "Actual_kd")] == "65"
    assert got[("PROJA_01_01", "5_SC_Injection_Details", 2, "Injected_Virus")] == "virus-b"
    assert got[("PROJA_01_01", "0a_Metadata", 1, "Sex")] == "F"
    assert ("PROJA_01_01", "0a_Metadata", 1, "Notes") not in got           # '?' is not a value
    assert got[("PROJA_01_03", "3b_Manual_Tray", 1, "Sex")] == "Female"    # sex from the scoring tab
    assert not any(k[0] == "PROJA_01_02" for k in got)                     # planned row skipped
    assert not any(k[0].startswith("PROJB") for k in got)                   # other cohort skipped
    assert any("PROJB_01_01" in m for m in w)                               # ...and said so


def test_odc_tab_header_found_below_notes(tmp_path):
    p = tmp_path / "K.xlsx"
    _odc_workbook(p)
    recs = ar.read_odc_tab(p, "ASPA_11", [])
    fields = {(r["subject_id"], r["field"]): r["value"] for r in recs}
    assert fields[("ASPA_11_01", "Organism_Strain")] == "strain-k"
    assert fields[("ASPA_11_01", "Surgery_Date")] == "2022-09-22"
    assert {r["subject_id"] for r in recs} == {"ASPA_11_01"}


def test_import_replaces_instead_of_piling_up(tmp_path):
    from mousedb.database import Database
    from mousedb.schema import AnimalRecord
    db = Database(tmp_path / "t.db")
    p = tmp_path / "PROJA_01.xlsx"
    _cnt_workbook(p)
    first = ar.sync_cohort(db, "PROJA_01", p)
    second = ar.sync_cohort(db, "PROJA_01", p)
    assert first["error"] is None and second["records"] == first["records"] > 0
    with db.session() as s:
        assert s.query(AnimalRecord).count() == first["records"]


def test_sync_never_raises(tmp_path):
    r = ar.sync_cohort(None, "PROJA_01", tmp_path / "missing.xlsx")
    assert r["error"] and r["records"] == 0
