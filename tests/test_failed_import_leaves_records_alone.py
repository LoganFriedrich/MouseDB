"""The sheet's two readers must agree on what one import did, and say each thing once.

WHY THIS FILE EXISTS
--------------------
One tracking-sheet import runs two readers over the same workbook: the importer
(the typed tables -- animals, weights, scores, surgeries) and then
animal_records (each animal's details, copied column for column, REPLACING the
cohort's previous copy). Three defects let them disagree, or let a person hear
about one problem twice -- or not at all:

1. A failed import still replaced the animal details. The importer saves a
   workbook whole or not at all: when it records an error it undoes the whole
   workbook. sheet_sync ran animal_records anyway, so after that run the animal
   details described the new workbook while every other table still held the
   old one. The animal details must now be left as they were, and the ledger
   line must say so and why.

2. An animal cell neither reader can read was named TWICE per import -- by the
   importer ("the animal cell is not an animal id") and by animal_records
   ("row with unrecognised animal ... skipped"). The fix must not overshoot:
   the importer reads a row's animal id under only some headings (on
   0a_Metadata, 'SubjectID' and 'Animal' only), and a cell under any other
   heading is read by animal_records alone. Staying silent there would leave
   the row named by NOBODY. So the tests below run the REAL importer and the
   REAL animal_records under every heading animal_records accepts, on every tab
   both read, and require exactly one line per bad cell -- from the importer
   where it reads that cell, from animal_records everywhere else. If the
   importer starts or stops reading a heading without
   animal_records.IMPORTER_ANIMAL_ID_HEADINGS following, a count turns 0 or 2.

3. manual_setup_check.py imported the sheets by calling the importer directly,
   bypassing the import ledger, the cross-task database lock and the one-sheet-
   per-cohort rule every other route uses. It must go through
   sheet_sync.import_cohorts, still dry run first, then the real import.
"""
import builtins
import contextlib
import importlib.util
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import mousedb.animal_records as ar
import mousedb.cohort_sheets as cs
import mousedb.database as dbmod
import mousedb.importers as importers
import mousedb.sheet_sync as ss
from mousedb.schema import AnimalRecord, Cohort, Subject

SHEET = "Connectome_05_Animal_Tracking.xlsx"


@pytest.fixture(autouse=True)
def no_production_side_effects(monkeypatch, tmp_path):
    """No test here may take the PRODUCTION database-task lock or write audit
    lines into the configured logs folder: the lock is a no-op, the audit log
    goes to tmp_path."""
    from mousedb import task_mutex
    monkeypatch.setattr(task_mutex, "hold", lambda *a, **k: contextlib.nullcontext())
    monkeypatch.setattr(dbmod, "DEFAULT_LOG_PATH", tmp_path / "audit_logs")


@pytest.fixture
def folder(tmp_path, monkeypatch):
    """A sheets folder holding cohort 05's sheet (contents written per test), a
    private config and a private ledger."""
    d = tmp_path / "sheets"
    d.mkdir()
    (d / SHEET).write_bytes(b"x")
    monkeypatch.setattr(cs, "CONFIG_PATH", tmp_path / "config.json")
    monkeypatch.setenv(cs.ENV_VAR, str(d))
    monkeypatch.delenv(cs.ASPA_ENV_VAR, raising=False)
    monkeypatch.setattr(ss, "LEDGER", tmp_path / "ledger.jsonl")
    return d


def _ledger():
    p = Path(ss.LEDGER)
    return [json.loads(l) for l in p.read_text().splitlines()] if p.exists() else []


# ---------------------------------------------------------------------------
# 1. an import that saved nothing leaves the animal details as they were
# ---------------------------------------------------------------------------

# What ExcelImporter returns when a reader recorded an error and the whole
# workbook was undone (see ExcelImporter._report_not_committed).
ROLLED_BACK = {"success": False, "imported": {"subjects": 0}, "not_saved": {"subjects": 16},
               "warnings": [], "errors": ["a reader failed",
                                          "Nothing from %s was saved: ..." % SHEET]}


@pytest.fixture
def copies(monkeypatch):
    """Every request to animal_records.sync_cohort, recorded instead of run."""
    calls = []

    def fake_sync(db, cohort_id, sheet, kind, dry_run=False):
        calls.append((cohort_id, kind, dry_run))
        return {"records": 7, "warnings": [], "error": None}

    monkeypatch.setattr(ar, "sync_cohort", fake_sync)
    return calls


def _importer_returning(monkeypatch, result):
    class Fake:
        db = "a database"          # so sheet_sync would copy the animal details

        def import_cohort_file(self, path, dry_run=False):
            if isinstance(result, Exception):
                raise result
            return dict(result)

    monkeypatch.setattr(importers, "ExcelImporter", lambda *a, **k: Fake())


@pytest.mark.parametrize("result", [
    ROLLED_BACK,
    {"success": False, "imported": {}, "warnings": [], "errors": ["Failed to open Excel file"]},
    {"success": False, "imported": {}, "warnings": [], "errors": []},  # failed, reason not given
], ids=["rolled_back", "error_before_reading", "failed_without_reason"])
def test_a_failed_import_does_not_replace_the_animal_details(folder, monkeypatch, copies, result):
    _importer_returning(monkeypatch, result)
    (c,) = ss.import_cohorts(["CNT_05"], triggered_by="test")["cohorts"]
    assert copies == []                                   # sync_cohort never asked
    assert c["success"] is False
    assert "animal_records" not in c["imported"]          # no count that reads as "copied 0"
    note = c[ss.LEFT_UNCHANGED]
    assert "left as they were" in note and "nothing from this workbook was saved" in note
    (line,) = _ledger()                                   # ...and the history says so
    assert line[ss.LEFT_UNCHANGED] == note and line["success"] is False


def test_a_saved_import_still_replaces_them(folder, monkeypatch, copies):
    _importer_returning(monkeypatch, {"success": True, "imported": {"subjects": 1},
                                      "warnings": [], "errors": []})
    (c,) = ss.import_cohorts(["CNT_05"], triggered_by="test")["cohorts"]
    assert copies == [("CNT_05", "cohort", False)]
    assert c["success"] and c["imported"]["animal_records"] == 7
    assert ss.LEFT_UNCHANGED not in c and ss.LEFT_UNCHANGED not in _ledger()[0]


def test_a_dry_run_previews_the_same_decision_and_records_nothing(folder, monkeypatch, copies):
    _importer_returning(monkeypatch, {"success": False, "imported": {}, "warnings": [],
                                      "errors": ["a reader failed"]})
    (c,) = ss.import_cohorts(["CNT_05"], dry_run=True)["cohorts"]
    assert copies == []
    assert "would be left as they are" in c[ss.LEFT_UNCHANGED]
    assert _ledger() == []


@pytest.mark.parametrize("hold_fails", [False, True], ids=["importer_crash", "lock_never_freed"])
def test_a_crash_before_the_copy_is_recorded_as_leaving_them_alone(folder, monkeypatch, copies,
                                                                   hold_fails):
    _importer_returning(monkeypatch, RuntimeError("workbook exploded"))
    if hold_fails:
        from mousedb import task_mutex

        def never(*a, **k):
            raise RuntimeError("could not acquire the database-task lock")
        monkeypatch.setattr(task_mutex, "hold", never)
    (c,) = ss.import_cohorts(["CNT_05"], triggered_by="test")["cohorts"]
    assert copies == [] and c["success"] is False
    assert "left as they were" in c[ss.LEFT_UNCHANGED]
    assert "left as they were" in _ledger()[0][ss.LEFT_UNCHANGED]


def test_the_console_says_it_under_the_fail_line(folder, monkeypatch, copies):
    _importer_returning(monkeypatch, ROLLED_BACK)
    r = ss.import_cohorts(["CNT_05"], triggered_by="test")
    lines = []
    assert ss.print_import_result(r, log=lines.append) is False
    fail = next(i for i, l in enumerate(lines) if l.startswith("FAIL CNT_05"))
    assert lines[fail + 1].startswith("       [!] animal details left as they were")
    assert lines[-1].startswith("Sheets:")                # the verdict stays the last line
    assert all(l.isascii() for l in lines)


def _workbook_05(path, sex):
    """A minimal real cohort-05 workbook: one animal on 0a_Metadata."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "0a_Metadata"
    ws.append(["SubjectID", "Sex"])
    ws.append(["CNT_05_01", sex])
    wb.save(path)


def test_end_to_end_a_rolled_back_workbook_leaves_the_old_details_in_place(folder, tmp_path,
                                                                          monkeypatch):
    """The real importer and the real animal_records over a scratch database:
    the database holds an older copy of the animal's details ('M'); the sheet
    now says 'F'. An import whose reader records an error saves nothing -- and
    the details must still read 'M'. The same workbook imported cleanly then
    replaces them."""
    real = importers.ExcelImporter
    db = dbmod.Database(tmp_path / "t.db")
    db.init_db()
    with db.session() as s:
        s.add(Cohort(cohort_id="CNT_05", project_code="CNT", start_date=date(2025, 6, 1)))
    ar.replace_cohort(db, "CNT_05", ar.CNT_TABS + (ar.TRAY_TAB,), [{
        "subject_id": "CNT_05_01", "cohort_id": "CNT_05", "source_tab": "0a_Metadata",
        "record_no": 1, "field": "Sex", "value": "M", "source_file": "older.xlsx"}])
    _workbook_05(folder / SHEET, "F")

    class FailsPartWay(importers.ExcelImporter):
        def _import_stats(self, *a, **k):
            self.errors.append("a reader recorded an error part way (test)")

    def sex_on_record():
        with db.session() as s:
            return [r.value for r in s.query(AnimalRecord).filter_by(
                subject_id="CNT_05_01", field="Sex")]

    monkeypatch.setattr(importers, "ExcelImporter", lambda *a, **k: FailsPartWay(db))
    (c,) = ss.import_cohorts(["CNT_05"], triggered_by="test")["cohorts"]
    assert c["success"] is False and "Nothing from" in c["error"]
    with db.session() as s:
        assert s.query(Subject).filter_by(subject_id="CNT_05_01").count() == 0  # undone
    assert sex_on_record() == ["M"]                       # left as they were
    assert "left as they were" in _ledger()[-1][ss.LEFT_UNCHANGED]

    monkeypatch.setattr(importers, "ExcelImporter", lambda *a, **k: real(db))
    (c,) = ss.import_cohorts(["CNT_05"], triggered_by="test")["cohorts"]
    assert c["success"] is True and ss.LEFT_UNCHANGED not in c
    assert sex_on_record() == ["F"]                       # now replaced


# ---------------------------------------------------------------------------
# 2. an unreadable animal cell is named once -- never twice, never by nobody
# ---------------------------------------------------------------------------

TODAY = date(2025, 10, 1)
DAY = datetime(2025, 7, 2)
SURGERY_TABS = ("4_Contusion_Injury_Details", "5_SC_Injection_Details")
# Cells both readers refuse: text that is no id, a number no animal can have,
# and another cohort's animal in both spellings.
REFUSED_BY_BOTH = ("mouse x", "123", "CNT_02_05", "CNT-02-07")


@pytest.fixture
def scratch_db(tmp_path):
    db = dbmod.Database(tmp_path / "t.db")
    db.init_db()
    with db.session() as s:
        s.add(Cohort(cohort_id="CNT_01", project_code="CNT", start_date=date(2025, 6, 1)))
    return db


def _tab_workbook(path, tab, header, ids):
    """One tab, ``header`` first; one row per id. On a surgery tab every row is
    a PERFORMED surgery (weighed and dated), the only rows whose id matters."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab
    surgery = tab in SURGERY_TABS
    ws.append(list(header) + (["Surgery_Date", "Subject_Weight (g)"] if surgery else ["Sex"]))
    for i in ids:
        cells = [i] if not isinstance(i, (list, tuple)) else list(i)
        ws.append(cells + ([DAY, 23.0] if surgery else ["M"]))
    wb.save(path)
    return path


def _both_readers(scratch_db, tmp_path, tab, header, ids):
    """(the importer's warnings, animal_records' warnings) for one workbook, run
    the way the sheet import runs them (importer first, then animal_records)."""
    p = _tab_workbook(tmp_path / "Connectome_01_Animal_Tracking.xlsx", tab, header, ids)
    imp = importers.ExcelImporter(scratch_db)
    imp.today = TODAY
    result = imp.import_cohort_file(p, dry_run=True)
    w = []
    ar.read_cnt_workbook(p, "CNT_01", w)
    return result["warnings"], w


def _naming(lines, cell):
    """Lines that name ``cell`` as a person would search for it (the importer
    quotes unreadable text and names an id by its normal spelling)."""
    return [l for l in lines if "'%s'" % cell in l or cell.replace("-", "_") in l]


@pytest.mark.parametrize("tab", ar.CNT_TABS[:1] + SURGERY_TABS)
@pytest.mark.parametrize("heading", ar.SUBJECT_COLUMNS)
def test_each_refused_cell_is_named_exactly_once_under_every_heading(scratch_db, tmp_path,
                                                                     tab, heading):
    by_importer, by_records = _both_readers(scratch_db, tmp_path, tab, [heading],
                                            ("CNT_01_01",) + REFUSED_BY_BOTH)
    importer_reads_it = heading in ar.IMPORTER_ANIMAL_ID_HEADINGS[tab]
    for cell in REFUSED_BY_BOTH:
        named_i, named_r = _naming(by_importer, cell), _naming(by_records, cell)
        assert len(named_i) + len(named_r) == 1, (cell, by_importer, by_records)
        # ...and by the right reader: the importer where it reads this heading,
        # animal_records everywhere else (no silent gap).
        assert bool(named_i) == importer_reads_it, (cell, by_importer, by_records)


def test_an_id_only_the_importer_can_read_is_named_by_animal_records(scratch_db, tmp_path):
    """'CNT0103' is imported by the importer but not copied here: animal_records'
    line is the only notice that the two readers took the row differently."""
    by_importer, by_records = _both_readers(scratch_db, tmp_path, "0a_Metadata",
                                            ["SubjectID"], ["CNT0103"])
    assert _naming(by_importer, "CNT0103") == []
    assert len(_naming(by_records, "CNT0103")) == 1


def test_text_the_importer_reads_as_blank_is_still_named(scratch_db, tmp_path):
    """pandas turns 'n/a' into a blank, which the importer skips without a word
    on 0a_Metadata: animal_records must name it."""
    by_importer, by_records = _both_readers(scratch_db, tmp_path, "0a_Metadata",
                                            ["SubjectID"], ["n/a"])
    assert _naming(by_importer, "n/a") == []
    assert len(_naming(by_records, "n/a")) == 1


@pytest.mark.parametrize("tab", SURGERY_TABS)
def test_on_a_surgery_tab_the_importer_names_that_row_itself(scratch_db, tmp_path, tab):
    """On a surgery tab the importer names a weighed row it reads as blank-id,
    by its sheet row; animal_records must not add a second line about it."""
    by_importer, by_records = _both_readers(scratch_db, tmp_path, tab, ["Subject_ID"],
                                            ["CNT_01_01", "n/a"])
    assert len([l for l in by_importer if "animal id is blank" in l and "sheet row 3" in l]) == 1
    assert _naming(by_records, "n/a") == []


def test_two_id_columns_each_cell_is_named_by_the_reader_that_reads_it(scratch_db, tmp_path):
    """animal_records reads 'Subject_ID' first, the importer's 0a reader reads
    'SubjectID': different cells, each named once, by its own reader."""
    by_importer, by_records = _both_readers(scratch_db, tmp_path, "0a_Metadata",
                                            ["Subject_ID", "SubjectID"],
                                            [("mouse a", "mouse b")])
    assert _naming(by_importer, "mouse a") == [] and len(_naming(by_records, "mouse a")) == 1
    assert len(_naming(by_importer, "mouse b")) == 1 and _naming(by_records, "mouse b") == []


def test_a_heading_with_a_trailing_space_is_not_the_importers(scratch_db, tmp_path):
    """The importer names its columns by the header cell as typed: 'SubjectID '
    is not 'SubjectID' to it, though animal_records reads it (it strips)."""
    by_importer, by_records = _both_readers(scratch_db, tmp_path, "0a_Metadata",
                                            ["SubjectID "], ["mouse x"])
    assert _naming(by_importer, "mouse x") == []
    assert len(_naming(by_records, "mouse x")) == 1


def test_a_planned_surgerys_id_is_not_named_by_either(scratch_db, tmp_path):
    """Unchanged policy: a plan's id is not a problem until the surgery happens."""
    p = tmp_path / "Connectome_01_Animal_Tracking.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SURGERY_TABS[0]
    ws.append(["Subject_ID", "Surgery_Date", "Subject_Weight (g)"])
    ws.append(["mouse x", None, None])
    wb.save(p)
    imp = importers.ExcelImporter(scratch_db)
    imp.today = TODAY
    shown = imp.import_cohort_file(p, dry_run=True)["warnings"]
    w = []
    ar.read_cnt_workbook(p, "CNT_01", w)
    assert _naming(shown + w, "mouse x") == []


@pytest.mark.parametrize("tab", ar.CNT_TABS[:1] + SURGERY_TABS)
def test_a_tick_box_among_animal_numbers_is_still_named(scratch_db, tmp_path, tab):
    """pandas reads a whole column at once, and among numbers a TRUE cell
    becomes 1: the importer files that row under animal 1 without a word.
    Judged on its own the cell stays TRUE, which the importer would refuse --
    so animal_records must not take the importer's silence as an announcement
    here, or the row is named by nobody."""
    heading = ar.IMPORTER_ANIMAL_ID_HEADINGS[tab][0]
    by_importer, by_records = _both_readers(scratch_db, tmp_path, tab, [heading],
                                            [1, True, 2])
    assert _naming(by_importer, "True") == []
    assert len(_naming(by_records, "True")) == 1


# ---------------------------------------------------------------------------
# 3. manual_setup_check.py imports through the recorded, locked route
# ---------------------------------------------------------------------------

def _load_manual_setup_check():
    path = Path(__file__).resolve().parents[1] / "manual_setup_check.py"
    spec = importlib.util.spec_from_file_location("manual_setup_check_under_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def setup_check(monkeypatch, tmp_path):
    """The script with everything it touches replaced: a fake database that
    holds no animals, one cohort sheet in the status listing, and every sheet
    import recorded instead of run. The old direct route fails the test."""
    import mousedb.exporters as exporters

    class FakeDB:
        db_path = tmp_path / "not_created.db"
        subjects = 0

        def get_stats(self):
            return {"subjects": FakeDB.subjects}

    class FakeExporter:
        def __init__(self, db):
            pass

        def __getattr__(self, name):
            return lambda *a, **k: self if name != "count" else 0

    imports = []

    def fake_import_cohorts(cohorts=None, dry_run=False, triggered_by="manual", sheets_dir=None):
        imports.append((cohorts, dry_run, triggered_by))
        return {"dry_run": dry_run, "cohorts": [], "triggered_by": triggered_by}

    def old_route(*a, **k):
        pytest.fail("manual_setup_check called importers.import_all_cohorts directly")

    monkeypatch.setattr(dbmod, "init_database", lambda *a, **k: FakeDB())
    monkeypatch.setattr(exporters, "QueryExporter", FakeExporter)
    monkeypatch.setattr(importers, "import_all_cohorts", old_route)
    monkeypatch.setattr(ss, "status", lambda: {"cohorts": [{
        "cohort_id": "CNT_05", "project": ss.SHEET_FOLDER_PROJECT, "sheet": SHEET}],
        "problem": None})
    monkeypatch.setattr(ss, "main", lambda argv=None: 0)
    monkeypatch.setattr(ss, "import_cohorts", fake_import_cohorts)
    return _load_manual_setup_check(), imports, FakeDB


def _answers(monkeypatch, *said):
    it = iter(said)

    def answer(*a):
        try:
            return next(it)
        except StopIteration:
            # pytest.fail, not StopIteration: step 4 reports an Exception and
            # carries on, which would hide a question the test did not expect.
            pytest.fail("manual_setup_check asked a question this test did not expect")

    monkeypatch.setattr(builtins, "input", answer)


def test_setup_check_dry_runs_then_imports_through_sheet_sync(setup_check, monkeypatch):
    script, imports, _ = setup_check
    _answers(monkeypatch, "y", "y")
    script.main()
    assert imports == [(None, True, "manual_setup_check"), (None, False, "manual_setup_check")]


def test_setup_check_stops_after_the_dry_run_when_told_no(setup_check, monkeypatch):
    script, imports, _ = setup_check
    _answers(monkeypatch, "y", "n")
    script.main()
    assert imports == [(None, True, "manual_setup_check")]


def test_setup_check_never_imports_into_a_database_that_holds_animals(setup_check, monkeypatch,
                                                                      capsys):
    script, imports, FakeDB = setup_check
    FakeDB.subjects = 5
    _answers(monkeypatch)                                  # any question asked fails
    script.main()
    assert imports == []
    assert "already holds animals" in capsys.readouterr().out


def test_setup_check_reports_a_sheet_step_failure_and_carries_on(setup_check, monkeypatch,
                                                                 capsys):
    """sheet_sync can raise where the old folder listing could not (its import
    history needs mousedb_root's logs folder). Like steps 2 and 3, step 4
    reports the error and the smoke check still runs steps 5 and 6."""
    script, imports, _ = setup_check

    def unconfigured():
        raise RuntimeError("mousedb does not know 'mousedb_root' on this machine")

    monkeypatch.setattr(ss, "status", unconfigured)
    _answers(monkeypatch)                                  # any question asked fails
    script.main()
    out = capsys.readouterr().out
    assert "Error: RuntimeError: mousedb does not know 'mousedb_root'" in out
    assert "[5] Testing QueryExporter" in out and "Setup test complete!" in out
    assert imports == []
