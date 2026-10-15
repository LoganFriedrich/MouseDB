"""A surgery row counts as a surgery that HAPPENED only when the weight on the day is typed in.

WHY these tests exist: the rule that decided whether a row of the injection tab
recorded an injection accepted Virus_Titer, Injected_Virus and Injection_Target
as proof. Those columns are filled in BEFORE the day -- the blank-sheet
generator and a workbook copied from the previous cohort carry them on every
row -- so a whole cohort's planned injections were copied into animal_records
as injections that had happened, and went out in the shared ODC dataset with a
titre and a 0 mL anaesthetic volume for animals nobody had operated on. The
contusion rule had the same hole (a dwell of 0 and Survived both outlive a
partial clearing of the row).

The proof is the one number that cannot exist before the day: the animal's
body weight typed on the row (it is weighed on the table to dose its drugs),
on a dated row. The surgery importer judges the same rows with the SAME rule
(animal_records.surgery_verdict); the agreement tests at the bottom fail if the
two readers ever take different surgeries from one sheet again -- they once
kept separate copies of the rule, and the copies drifted.

Rows left out are announced ONCE per import, by the importer only (it runs on
every route that runs this module): a held row -- one that looks done but
cannot be taken as done -- every time, with what to fix; a planned row only
once it is overdue (dated more than PLANNED_SURGERY_GRACE_DAYS ago), because
a plan is not a problem: the sheet is filled in ahead of time on purpose.
"""
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd
import pytest

from mousedb import animal_records as ar

CONT, INJ = "4_Contusion_Injury_Details", "5_SC_Injection_Details"
WEIGHT = "Subject_Weight (g)"

# Headers as the live workbooks write them (row 1 of each surgery tab).
CONT_HEADER = ["Subject_ID", "Surgery_Date", "Surgery_Type", "Surgery_Severity",
               "Contusion_Location", WEIGHT, "Anesthetic", "Anesthetic_Dose",
               "Anesthetic_Volume", "Intended_kd", "Intended_Dwell", "Stage_Height",
               "Actual_kd", "Actual_displacement", "Actual_Velocity", "Actual_Dwell",
               "Survived"]
INJ_HEADER = ["Subject_ID", "Surgery_Date", WEIGHT, "Surgery_Type", "Injected_Virus",
              "Virus_Titer", "Injection_Target", "Depths (D/V)", "Coordinates (M/L)",
              "Anesthetic", "Anesthetic_Dose", "Anesthetic_Volume", "Survived"]

D1, D2 = datetime(2025, 7, 2), datetime(2025, 8, 20)
# The day the grace window is counted from in the tests that run the importer.
TODAY = date(2025, 10, 1)


def _cont(sid, date=None, weight=None, kd=None, disp=None, dwell=0, survived=None):
    """One contusion row. Defaults are what a planned row carries: the protocol
    constants, a Stage_Height and dwell of 0, a volume formula that evaluated to 0."""
    return [sid, date, "Contusion", "severity-a", "level-a", weight, "drug-x", "10mg/kg", 0,
            50, 0, 0, kd, disp, None, dwell, survived]


def _inj(sid, date=None, weight=None, survived="Y"):
    """One injection row. Every column a planned row is pre-filled with is filled
    in -- virus, titre, target, depths, coordinates, drugs, the 0 volume and even
    Survived -- so only the weight and the date tell a plan from a surgery."""
    return [sid, date, weight, "Spinal Injection", "virus-a", "1x10^9", "target-a",
            "-1.0/-0.5", "+/-0.5", "drug-x", "10mg/kg", 0, survived]


def _workbook(path, cont_rows, inj_rows, cont_header=CONT_HEADER, inj_header=INJ_HEADER,
              meta_id="PROJA_01_01"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CONT
    ws.append(cont_header)
    for r in cont_rows:
        ws.append(r)
    sc = wb.create_sheet(INJ)
    sc.append(inj_header)
    for r in inj_rows:
        sc.append(r)
    meta = wb.create_sheet("0a_Metadata")
    meta.append(["SubjectID", "Sex"])
    meta.append([meta_id, "F"])
    wb.save(path)
    return path


def _read(path):
    w = []
    recs = ar.read_cnt_workbook(path, "PROJA_01", w)
    return recs, w


def _copied(recs, tab):
    return {r["subject_id"] for r in recs if r["source_tab"] == tab}


# --- the defect itself ---------------------------------------------------------------

def test_planned_injection_with_prefilled_virus_and_titre_is_not_copied(tmp_path):
    p = _workbook(tmp_path / "PROJA_01.xlsx", [],
                  [_inj("PROJA_01_01", D2, 24.5),     # performed: weighed and dated
                   _inj("PROJA_01_02"),               # planned: everything but weight and date
                   _inj("PROJA_01_03")])
    recs, w = _read(p)
    assert _copied(recs, INJ) == {"PROJA_01_01"}
    done = {r["field"]: r["value"] for r in recs if r["subject_id"] == "PROJA_01_01"
            and r["source_tab"] == INJ}
    # A performed row is still copied column for column, titre included.
    assert done["Virus_Titer"] == "1x10^9" and done[WEIGHT] == "24.5"
    # Not announced here: the importer is the one place (see the module docstring).
    assert w == []


def test_contusion_needs_the_weight_not_a_zero_dwell_or_survived(tmp_path):
    # A dated row with Survived and a dwell of 0 but no weight: what a partially
    # cleared copy of last cohort's row, or a generator date, looks like.
    p = _workbook(tmp_path / "PROJA_01.xlsx",
                  [_cont("PROJA_01_01", D1, 25.5, kd=65, disp=500, survived="Y"),
                   _cont("PROJA_01_02", D1, None, survived="Y")], [])
    recs, w = _read(p)
    assert _copied(recs, CONT) == {"PROJA_01_01"}
    assert w == []


# --- rows that look done but cannot be taken as done are held, not copied -------------

@pytest.mark.parametrize("cont_rows, inj_rows", [
    ([_cont("PROJA_01_01", D1, None, kd=65, disp=500)], []),       # reading, no weight
    ([], [_inj("PROJA_01_01", None, 22)]),                          # weight, no date
    ([], [_inj("PROJA_01_01", "TBD", 22)]),                         # date cell not a date
    ([], [_inj("PROJA_01_01", D2, "22 \u00b5g")]),                  # not a weight in grams
])
def test_a_held_row_is_not_copied(tmp_path, cont_rows, inj_rows):
    recs, w = _read(_workbook(tmp_path / "PROJA_01.xlsx", cont_rows, inj_rows))
    assert _copied(recs, CONT) == set() and _copied(recs, INJ) == set()
    assert w == []


@pytest.mark.parametrize("when, verdict", [
    ("2025-07-02", ar.PERFORMED), ("2025-07-02 13:00:00", ar.PERFORMED),
    ("8/27/2026", ar.PERFORMED), (datetime(2025, 7, 2), ar.PERFORMED),
    (pd.Timestamp("2025-07-02"), ar.PERFORMED), (date(2025, 7, 2), ar.PERFORMED),
    (None, ar.HELD), ("", ar.HELD), (float("nan"), ar.HELD), (pd.NaT, ar.HELD),
    ("TBD", ar.HELD), ("27\u201308", ar.HELD), ("NaT", ar.HELD),
    (45870, ar.HELD),       # a bare number: pandas would make it a 1970 timestamp
    ("today", ar.HELD),     # parses to whatever day the import runs: invented
    # Text that leaves part of the day out: pandas fills the gap in, from the
    # clock for a time ('13:00' becomes today at 13:00) and with the 1st for a
    # month or a year. Each would be a surgery day nobody recorded.
    ("13:00", ar.HELD), ("10:30 AM", ar.HELD), ("Aug 2025", ar.HELD), ("2025", ar.HELD),
    ("20250827", ar.PERFORMED), ("27 Aug 2025", ar.PERFORMED),
])
def test_what_counts_as_a_date_on_a_weighed_row(when, verdict):
    assert ar.surgery_verdict({"Surgery_Date": when, WEIGHT: "22"}, INJ)[0] == verdict


def test_tab_without_a_weight_column_copies_nothing(tmp_path):
    header = [h for h in INJ_HEADER if h != WEIGHT]
    row = [v for h, v in zip(INJ_HEADER, _inj("PROJA_01_01", D2, 24.5)) if h != WEIGHT]
    p = _workbook(tmp_path / "PROJA_01.xlsx", [], [row], inj_header=header)
    recs, w = _read(p)
    assert _copied(recs, INJ) == set()
    assert w == []      # the importer tells the person to restore the heading
    # The other tabs of the workbook are unaffected.
    assert _copied(recs, "0a_Metadata") == {"PROJA_01_01"}


def test_a_hand_made_weight_heading_is_read(tmp_path):
    # A hand-made sheet may head the weight column differently; the importer
    # accepts the same headings, so both must read it.
    header = ["Weight (g)" if h == WEIGHT else h for h in INJ_HEADER]
    p = _workbook(tmp_path / "PROJA_01.xlsx", [],
                  [_inj("PROJA_01_01", D2, 24.5), _inj("PROJA_01_02")], inj_header=header)
    recs, _ = _read(p)
    assert _copied(recs, INJ) == {"PROJA_01_01"}


def test_a_row_holding_only_the_animal_id_is_not_announced(tmp_path):
    p = _workbook(tmp_path / "PROJA_01.xlsx", [], [["PROJA_01_04"]])
    recs, w = _read(p)
    assert w == []


# --- the rule itself, on single rows -------------------------------------------------

@pytest.mark.parametrize("weight, verdict", [
    ("22", ar.PERFORMED), ("25.5", ar.PERFORMED), ("25,5", ar.PERFORMED),
    ("22 g", ar.PERFORMED), ("22g", ar.PERFORMED), ("22 grams", ar.PERFORMED),
    (25.5, ar.PERFORMED), (22, ar.PERFORMED),
    (None, ar.PLANNED), ("", ar.PLANNED), ("?", ar.PLANNED), ("0", ar.PLANNED),
    (0, ar.PLANNED), (float("nan"), ar.PLANNED),
    ("abc", ar.HELD), ("22 kg", ar.HELD), ("-3", ar.HELD), (-3, ar.HELD),
    ("1e1", ar.HELD), ("inf", ar.HELD),
    (True, ar.HELD), (datetime(2025, 7, 2), ar.HELD),
])
def test_what_counts_as_a_weight(weight, verdict):
    row = {"Surgery_Date": "2025-07-02", WEIGHT: weight}
    assert ar.surgery_verdict(row, INJ)[0] == verdict
    assert ar.surgery_was_performed(row, INJ) == (verdict == ar.PERFORMED)


@pytest.mark.parametrize("field", ["Injected_Virus", "Virus_Titer", "Injection_Target",
                                   "Survived", "Depths (D/V)", "Coordinates (M/L)"])
def test_no_prefillable_injection_column_is_proof(field):
    assert ar.surgery_verdict({"Surgery_Date": "2025-08-20", field: "x"}, INJ)[0] == ar.PLANNED


@pytest.mark.parametrize("field", ["Actual_Dwell", "Survived", "Actual_Velocity",
                                   "Intended_kd", "Stage_Height"])
def test_no_prefillable_contusion_column_is_proof(field):
    assert ar.surgery_verdict({"Surgery_Date": "2025-07-02", field: "50"}, CONT)[0] == ar.PLANNED


def test_a_tab_with_no_rule_is_refused_not_let_through():
    with pytest.raises(KeyError):
        ar.surgery_verdict({WEIGHT: "22"}, "6_Ladder")


def test_the_first_filled_heading_decides_weight_and_date():
    # A blank first weight heading beside a weight under a later one, and a
    # blank Surgery_Date beside a filled Date: the filled cell is read.
    row = {WEIGHT: None, "Weight (g)": "23", "Surgery_Date": float("nan"), "Date": "2025-08-20"}
    assert ar.surgery_verdict(row, INJ)[0] == ar.PERFORMED
    assert ar.surgery_weight(row, INJ) == 23.0
    assert ar.surgery_date(row) == date(2025, 8, 20)


@pytest.mark.parametrize("days_ago, overdue", [
    (None, False),                                    # undated: never announced
    (-5, False),                                      # still ahead
    (0, False),
    (ar.PLANNED_SURGERY_GRACE_DAYS, False),           # on the edge: still in grace
    (ar.PLANNED_SURGERY_GRACE_DAYS + 1, True),        # overdue
])
def test_a_plan_is_overdue_only_after_the_grace_window(days_ago, overdue):
    when = None if days_ago is None else TODAY - timedelta(days=days_ago)
    row = {"Surgery_Date": when, WEIGHT: None}
    assert ar.planned_surgery_overdue(row, INJ, TODAY) == (when if overdue else None)
    # A performed or held row is never "an overdue plan", whatever its date.
    old = TODAY - timedelta(days=100)
    assert ar.planned_surgery_overdue({"Surgery_Date": old, WEIGHT: 22}, INJ, TODAY) is None
    assert ar.planned_surgery_overdue({"Surgery_Date": old, WEIGHT: "abc"}, INJ, TODAY) is None


def test_rule_reads_pandas_rows_the_way_the_importer_holds_them(tmp_path):
    # The importer reads rows with pandas ('' or NaN blanks, datetimes, floats);
    # this module reads the workbook directly. The one rule must give the same
    # answer on both: on the rows the importer's gate judges (cells as typed)
    # and, for rows pandas does not convert, on pandas' ordinary read too.
    from mousedb.importers import ExcelImporter
    p = _workbook(tmp_path / "PROJA_01.xlsx",
                  [_cont("PROJA_01_01", D1, 25.5, kd=65, disp=500, survived="Y"),
                   _cont("PROJA_01_02", D1, None, survived="Y"),
                   _cont("PROJA_01_03", D1, None, kd=65, disp=500)],
                  [_inj("PROJA_01_01", D2, 24.5), _inj("PROJA_01_02"),
                   _inj("PROJA_01_03", D2, "23 g"), _inj("PROJA_01_04", None, 22),
                   _inj("PROJA_01_05", 2025, 22)])     # a number where the date goes
    recs, _ = _read(p)
    imp = ExcelImporter.__new__(ExcelImporter)
    for tab in (CONT, INJ):
        df = pd.read_excel(p, sheet_name=tab)
        by_pandas = {row["Subject_ID"] for _, row in df.iterrows()
                     if ar.surgery_was_performed(row, tab)}
        as_typed = {cells["Subject_ID"] for cells in imp._surgery_rows_as_typed(p, tab).values()
                    if ar.surgery_was_performed(cells, tab)}
        assert by_pandas == as_typed == _copied(recs, tab), tab


# --- existing data: the next import retires what the old rule stored ------------------

def test_next_sync_removes_planned_rows_the_old_rule_stored(tmp_path):
    from mousedb.database import Database
    from mousedb.schema import AnimalRecord
    db = Database(tmp_path / "t.db")
    # What the old rule left in the table: a planned injection copied as done.
    stale = [{"subject_id": "PROJA_01_02", "cohort_id": "PROJA_01", "source_tab": INJ,
              "record_no": 1, "field": f, "value": v, "source_file": "PROJA_01.xlsx"}
             for f, v in (("Virus_Titer", "1x10^9"), ("Anesthetic_Volume", "0"))]
    ar.replace_cohort(db, "PROJA_01", (INJ,), stale)
    p = _workbook(tmp_path / "PROJA_01.xlsx", [],
                  [_inj("PROJA_01_01", D2, 24.5), _inj("PROJA_01_02")])
    r = ar.sync_cohort(db, "PROJA_01", p)
    assert r["error"] is None
    assert r["warnings"] == []
    with db.session() as s:
        left = {(a.subject_id, a.source_tab) for a in s.query(AnimalRecord)}
    assert ("PROJA_01_02", INJ) not in left
    assert ("PROJA_01_01", INJ) in left and ("PROJA_01_01", "0a_Metadata") in left


# --- one import, both readers: each row is announced once, and only if it is a problem --

def _cnt_db(tmp_path):
    """A scratch database that already knows the cohort, so the importer's
    warnings are about the surgery rows only (no placeholder start)."""
    from mousedb import database as dbmod
    from mousedb.schema import Cohort
    db = dbmod.Database(tmp_path / "t.db")
    db.init_db()
    with db.session() as s:
        s.add(Cohort(cohort_id="CNT_01", project_code="CNT", start_date=date(2025, 6, 1)))
    return db


def _sheet_import(tmp_path, cont_rows, inj_rows, cont_header=CONT_HEADER,
                  inj_header=INJ_HEADER, dry_run=True):
    """Run the two readers of one workbook the way the sheet import does
    (sheet_sync._import_one: the importer, then animal_records) and return
    (every warning a person is shown, the importer's result, the records,
    the database)."""
    from mousedb.importers import ExcelImporter
    db = _cnt_db(tmp_path)
    p = _workbook(tmp_path / "Connectome_01_Animal_Tracking.xlsx", cont_rows, inj_rows,
                  cont_header=cont_header, inj_header=inj_header, meta_id="CNT_01_01")
    imp = ExcelImporter(db)
    imp.today = TODAY
    result = imp.import_cohort_file(p, dry_run=dry_run)
    rec = ar.sync_cohort(db, "CNT_01", p, dry_run=True)
    assert rec["error"] is None
    recs = ar.read_cnt_workbook(p, "CNT_01", [])
    return result["warnings"] + rec["warnings"], result, recs, db


def _lines_naming(warnings, animal):
    return [w for w in warnings if animal in w]


def test_plans_are_silent_until_overdue_and_then_announced_once(tmp_path):
    overdue = TODAY - timedelta(days=ar.PLANNED_SURGERY_GRACE_DAYS + 1)
    recent = TODAY - timedelta(days=ar.PLANNED_SURGERY_GRACE_DAYS)
    shown, _, _, _ = _sheet_import(
        tmp_path,
        [_cont("CNT_01_01"),                               # undated plan
         _cont("CNT_01_02", recent),                       # inside the grace window
         _cont("CNT_01_03", TODAY + timedelta(days=4)),    # still ahead
         _cont("CNT_01_04", overdue)],                     # overdue
        [_inj("CNT_01_05")])                               # undated injection plan
    for quiet in ("CNT_01_01", "CNT_01_02", "CNT_01_03", "CNT_01_05"):
        assert _lines_naming(shown, quiet) == [], quiet
    line = _lines_naming(shown, "CNT_01_04")
    assert len(line) == 1 and "planned but not recorded" in line[0]
    assert "CNT_01_04 (%s)" % overdue in line[0]


@pytest.mark.parametrize("cont_rows, inj_rows, needle", [
    ([_cont("CNT_01_01", D1, None, kd=65, disp=500)], [], "device reading"),
    ([], [_inj("CNT_01_01", None, 22)], "no Surgery_Date"),
    ([], [_inj("CNT_01_01", "TBD", 22)], "not a date"),
    ([], [_inj("CNT_01_01", D2, "22 \u00b5g")], "not a weight in grams"),
    ([], [_inj("CNT_01_01", "27\u201308", 22)], "not a date"),
])
def test_a_held_row_is_announced_once_in_plain_ascii(tmp_path, cont_rows, inj_rows, needle):
    shown, _, recs, _ = _sheet_import(tmp_path, cont_rows, inj_rows)
    line = _lines_naming(shown, "CNT_01_01")
    assert len(line) == 1 and needle in line[0], shown
    assert all(w.isascii() for w in shown)
    assert {r["source_tab"] for r in recs} == {"0a_Metadata"}


def test_another_cohorts_animal_is_announced_once_and_only_when_it_matters(tmp_path):
    # A performed surgery of a borrowed animal: both readers refuse it, and it
    # is named once (by the importer). A PLAN naming a borrowed animal, or a
    # skeleton row whose id was never finished, is not a problem: silent.
    shown, _, recs, _ = _sheet_import(
        tmp_path,
        [_cont("CNT_02_05", D1, 22.0, kd=60, disp=500, survived="Y"),   # performed, borrowed
         _cont("CNT_02_06"),                                            # plan, borrowed
         _cont("CNT_01_")],                                             # unfinished id, plan
        [])
    assert len(_lines_naming(shown, "CNT_02_05")) == 1
    assert _lines_naming(shown, "CNT_02_06") == []
    assert _lines_naming(shown, "CNT_01_'") == [] and _lines_naming(shown, "'CNT_01_") == []
    assert {r["source_tab"] for r in recs} == {"0a_Metadata"}


def test_a_missing_weight_column_is_announced_once(tmp_path):
    header = [h for h in INJ_HEADER if h != WEIGHT]
    rows = [[v for h, v in zip(INJ_HEADER, _inj(a, D2, 24.5)) if h != WEIGHT]
            for a in ("CNT_01_01", "CNT_01_02")]
    shown, _, _, _ = _sheet_import(tmp_path, [], rows, inj_header=header)
    line = [w for w in shown if INJ in w and "'%s'" % WEIGHT in w]
    assert len(line) == 1 and "restore" in line[0]
    assert _lines_naming(shown, "CNT_01_02") == []


# --- the importer and this module must never disagree again --------------------------

@pytest.mark.parametrize("heading", ar.SURGERY_WEIGHT_COLUMNS)
def test_importer_surgery_gate_agrees_with_the_shared_rule(tmp_path, heading):
    """Every row is judged the same way by the surgery importer and here, under
    every heading the weight may be typed under.

    The importer's side is ExcelImporter._performed_surgery: the gate BOTH of its
    surgery-tab readers call, row by row, before anything is stored. Comparing a
    helper the readers do not call would pass while the real gate drifted. For a
    row of this workbook's own cohort it answers only "did this surgery happen",
    which is exactly surgery_was_performed. A disagreement means the typed
    surgeries table and animal_records hold different surgeries for the same
    sheet -- the importer must decide with animal_records.surgery_verdict, not
    keep a copy of the rule. The rows include the inputs on which two copies of
    the rule once disagreed ('22 grams', '1e1', 'inf')."""
    from mousedb.database import Database
    from mousedb.importers import ExcelImporter
    imp = ExcelImporter(db=Database(tmp_path / "t.db"))
    p = _workbook(tmp_path / "PROJA_01.xlsx",
                  [_cont("PROJA_01_01", D1, 25.5, kd=65, disp=500, survived="Y"),
                   _cont("PROJA_01_02", D1, None, survived="Y"),           # dwell 0 + Survived
                   _cont("PROJA_01_03", D1, None, kd=65, disp=500),        # reading, no weight
                   _cont("PROJA_01_04", D1, 0),                            # weight 0
                   _cont("PROJA_01_05", D1, "22 g", kd=61, disp=480),      # unit typed in
                   _cont("PROJA_01_06", D1, "25,5", kd=61, disp=480),      # decimal comma
                   _cont("PROJA_01_07", None, 25.5, kd=61, disp=480),      # weighed, no date
                   _cont("PROJA_01_08"),                                   # skeleton, no date
                   _cont("PROJA_01_09", D1, "22 grams", kd=61, disp=480),  # unit as a word
                   _cont("PROJA_01_10", 45870, 22, kd=61, disp=480)],      # a number as the date
                  [_inj("PROJA_01_01", D2, 24.5),
                   _inj("PROJA_01_02", D2),                                # titre, virus, target...
                   _inj("PROJA_01_03", D2, "abc"),                         # unreadable weight
                   _inj("PROJA_01_04", D2, 0),
                   _inj("PROJA_01_05", D2, -3),                            # negative
                   _inj("PROJA_01_06", D2, "22 kg"),                       # wrong unit
                   _inj("PROJA_01_07", D2, "?"),                           # 'not recorded'
                   _inj("PROJA_01_08", None, 22),                          # weighed, no date
                   _inj("PROJA_01_09"),                                    # skeleton, no date
                   _inj("PROJA_01_10", "TBD", 23),                         # weighed, date not a date
                   _inj("PROJA_01_11", D2, "1e1"),                         # float() reads it
                   _inj("PROJA_01_12", D2, "inf")],                        # float() reads it
                  cont_header=[heading if h == WEIGHT else h for h in CONT_HEADER],
                  inj_header=[heading if h == WEIGHT else h for h in INJ_HEADER])
    disagree, judged = [], 0
    for tab in ar.SURGERY_TABS.values():
        # The rows exactly as the importer's surgery readers hand them to its
        # gate (cells as typed: see ExcelImporter._surgery_rows_as_typed).
        for cells in imp._surgery_rows_as_typed(p, tab).values():
            sid = cells["Subject_ID"]
            mine = ar.surgery_was_performed(cells, tab)
            theirs = imp._performed_surgery(cells, sid, "PROJA_01", tab) is not None
            judged += 1
            if mine != theirs:
                disagree.append((tab, sid, "animal_records=%s importer=%s" % (mine, theirs)))
    assert judged == 22           # every row above reached both gates
    assert disagree == []


def test_a_surgery_time_typed_as_the_date_is_not_stored_as_the_day_of_the_import(tmp_path):
    # Both readers would otherwise store the day the import ran: the importer
    # once (it only adds rows), animal_records a new day on every run.
    shown, result, recs, db = _sheet_import(
        tmp_path, [], [_inj("CNT_01_01", "13:00", 24.5)], dry_run=False)
    from mousedb.schema import Surgery
    with db.session() as s:
        assert s.query(Surgery).count() == 0
    assert {r["source_tab"] for r in recs} == {"0a_Metadata"}
    line = _lines_naming(shown, "CNT_01_01")
    assert len(line) == 1 and "not a date" in line[0], shown


def test_the_contusion_reader_judges_the_cells_as_typed_too(tmp_path):
    # The weight column holds only numbers and number-like text, so pandas'
    # ordinary read turns '1e1' into 10.0 and 'inf' into infinity. Judged on
    # that read, the contusion reader would store two surgeries animal_records
    # holds (each proved by a weight nobody typed). Both readers must hold the
    # same single surgery and name the other two rows as held.
    from mousedb.schema import Surgery
    header = ["Subject_ID", "Surgery_Date", WEIGHT, "Actual_kd", "Survived"]
    rows = [["CNT_01_01", D1, 25.5, 65, "Y"],
            ["CNT_01_02", D1, "1e1", 61, "Y"],
            ["CNT_01_03", D1, "inf", 60, "Y"]]
    assert pd.api.types.is_float_dtype(pd.read_excel(
        _workbook(tmp_path / "probe.xlsx", rows, [], cont_header=header),
        sheet_name=CONT)[WEIGHT])      # the premise: the ordinary read converts them
    shown, result, recs, db = _sheet_import(tmp_path, rows, [], cont_header=header,
                                            dry_run=False)
    assert result["success"], result["errors"]
    with db.session() as s:
        stored = {x.subject_id for x in s.query(Surgery).all()}
    assert stored == _copied(recs, CONT) == {"CNT_01_01"}
    for held in ("CNT_01_02", "CNT_01_03"):
        assert len(_lines_naming(shown, held)) == 1, (held, shown)


def test_headings_that_differ_only_by_spaces_are_judged_by_the_first_column(tmp_path):
    # animal_records strips headings and keeps the first column under a name;
    # the importer must judge the same column, or one reader takes a surgery
    # the other refuses.
    from mousedb.schema import Surgery
    header = ["Subject_ID", "Surgery_Date", WEIGHT, WEIGHT + " ", "Survived"]
    rows = [["CNT_01_01", D1, None, 25.5, "Y"],      # first column blank: a plan
            ["CNT_01_02", D1, 24.0, None, "Y"]]      # first column weighed: performed
    shown, result, recs, db = _sheet_import(tmp_path, rows, [], cont_header=header,
                                            dry_run=False)
    with db.session() as s:
        stored = {x.subject_id for x in s.query(Surgery).all()}
    assert stored == _copied(recs, CONT) == {"CNT_01_02"}


def test_both_readers_store_the_same_surgeries_from_one_sheet(tmp_path):
    """End to end: a real import and a real copy of ONE workbook hold the same
    surgeries -- same animal, same kind, same day.

    The sheet uses every way a cell can be read differently by two readers: a
    blank first weight heading with the weight under a later one, a blank
    Surgery_Date column with the day under Date, a heading typed with a
    trailing space, a weight with its unit as a word, a weight float() reads but
    the rule refuses ('1e1'), also in a column pandas would turn into numbers,
    and a number where the date goes. If either reader stops deciding with the
    shared rule, or stops judging the cells as typed, the two sets differ."""
    from mousedb.schema import Surgery
    cont_header = ["Subject_ID", "Surgery_Date", "Date", WEIGHT, "Weight (g)",
                   "Actual_kd ", "Survived"]
    inj_header = ["Subject_ID", "Surgery_Date ", WEIGHT + " ", "Injected_Virus", "Survived"]
    cont_rows = [
        ["CNT_01_01", D1, None, 25.5, None, 65, "Y"],          # plain
        ["CNT_01_02", None, D1, None, 23.0, 61, "Y"],          # second weight heading, Date column
        ["CNT_01_03", D1, None, "22 grams", None, 60, "Y"],    # unit as a word
        ["CNT_01_04", D1, None, "1e1", None, 60, "Y"],         # held
        ["CNT_01_05", 2025, None, 24.0, None, 60, "Y"],        # a number where the date goes
        ["CNT_01_06", D1, None, None, None, None, None],       # plan
    ]
    inj_rows = [
        ["CNT_01_01", D2, 24.5, "virus-a", "Y"],               # headings with a trailing space
        # Held. In a column whose every cell can be a number, pandas' ordinary
        # read turns this text into 10.0 -- a 10 g surgery to a reader that
        # judged that read.
        ["CNT_01_02", D2, "1e1", "virus-a", "Y"],
        ["CNT_01_03", D2, None, "virus-a", "Y"],               # plan with its virus pre-filled
    ]
    _, result, recs, db = _sheet_import(tmp_path, cont_rows, inj_rows, cont_header=cont_header,
                                        inj_header=inj_header, dry_run=False)
    assert result["success"], result["errors"]
    kind = {tab: k for k, tab in ar.SURGERY_TABS.items()}
    copied = set()
    by_record = {}
    for r in recs:
        if r["source_tab"] in kind:
            by_record.setdefault((r["subject_id"], r["source_tab"], r["record_no"]), {})[
                r["field"]] = r["value"]
    for (sid, tab, _n), fields in by_record.items():
        day = next(fields[f] for f in ar.SURGERY_DATE_FIELDS if fields.get(f))
        copied.add((sid, kind[tab], day))
    with db.session() as s:
        stored = {(x.subject_id, x.surgery_type, x.surgery_date.isoformat())
                  for x in s.query(Surgery).all()}
    assert stored == copied
    assert stored == {("CNT_01_01", "contusion", "2025-07-02"),
                      ("CNT_01_02", "contusion", "2025-07-02"),
                      ("CNT_01_03", "contusion", "2025-07-02"),
                      ("CNT_01_01", "tracing", "2025-08-20")}
