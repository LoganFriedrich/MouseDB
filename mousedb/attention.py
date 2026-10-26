"""What needs a person, gathered from every tool: shown in red on the Where Is
My Data tab, and written per cohort into a FIX THESE file beside its sheet.

WHY
---
A session whose hand scores cannot be found, a sheet whose date had to be
read as a typo, an import that failed: each is something only a person can
fix, usually in the tracking sheet. Left in a log it is never seen. So it is
put where people look -- the tab they open to find their data, and the folder
where the sheet itself is edited.

WHAT IS GATHERED
----------------
1. The reaching pipeline's own lists (<mousereach_pipeline_root>/attention/,
   items marked needs_person): today, processed videos whose session has no
   hand scores although the cohort's sheet has been imported since they were
   ready. The pipeline holds such a video unfiled until the scores come.
2. The tracking-sheet imports (the import ledger): for each cohort in the
   sheets folder, its last import if it failed, and every date that import
   had to repair (its REPAIRED warnings) -- correct those in the sheet.
3. Videos already in the database (Pillar trays) of the project whose sheets
   are kept current, whose session has no hand scores: filed before the
   pipeline checked hand scores at filing (2026-10-26), or named for an
   animal, date or tray the sheet does not have. Read from the analysis
   snapshot (so as of its last refresh) against the hand-scores file.
Frozen workbooks (the letter cohorts) are not listed: nobody corrects them.

THE FIX THESE FILE
------------------
After every import that writes, each cohort in the sheets folder with
anything listed gets "!! <cohort> - FIX THESE.xlsx" beside its workbook: one
line per thing to fix, what is wrong and what to do. A cohort with nothing
listed has the file removed. It is the one file mousedb writes into the
sheets folder besides a new cohort's blank sheet (cohort_sheets says why);
its name can never be read as a sheet.
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

FIX_THESE_NAME = "!! %s - FIX THESE.xlsx"


def mousereach_items(pipeline_root: Optional[Path] = None) -> List[dict]:
    """Items a person must act on from the reaching pipeline's lists."""
    if pipeline_root is None:
        from .config import mousereach_pipeline_root
        pipeline_root = mousereach_pipeline_root()
    out: List[dict] = []
    d = Path(pipeline_root) / "attention" if pipeline_root else None
    if not d or not d.is_dir():
        return out
    for f in sorted(d.glob("*.json")):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
        except Exception as e:
            out.append({"cohort_id": None, "source": "reaching pipeline",
                        "says": "The pipeline's list %s could not be read (%s)." % (f.name, e)})
            continue
        for i in data.get("items") or []:
            if i.get("needs_person"):
                out.append({"cohort_id": i.get("cohort_id"), "source": "reaching pipeline",
                            "says": i.get("says") or str(i)})
    return out


def sheet_items(cohort_ids: Optional[List[str]] = None) -> List[dict]:
    """From each cohort's last sheet import: a failure, and every REPAIRED
    date. ``cohort_ids``: the cohorts to look at (default: those in the
    sheets folder)."""
    from .sheet_sync import _read_ledger
    if cohort_ids is None:
        from .cohort_sheets import available_cohorts
        from .sheet_sync import _folder_cohort_id
        cohort_ids = [_folder_cohort_id(n) for n in available_cohorts()]
    last: Dict[str, dict] = {}
    for e in _read_ledger():
        if e.get("cohort_id") in cohort_ids and not e.get("dry_run"):
            last[e["cohort_id"]] = e
    out = []
    for cid in cohort_ids:
        e = last.get(cid)
        if not e:
            continue
        if not e.get("success"):
            out.append({"cohort_id": cid, "source": "sheet import",
                        "says": "The last import of %s's sheet (%s) failed: %s" % (
                            cid, e.get("finished") or e.get("started") or "?",
                            e.get("error") or "no reason recorded")})
        for w in e.get("warnings") or []:
            if str(w).startswith("REPAIRED"):
                out.append({"cohort_id": cid, "source": "sheet import", "says": str(w)})
    return out


def filed_without_scores(snapshot_dir: Optional[Path] = None,
                         hand_csv: Optional[Path] = None) -> List[dict]:
    """Source 3 above: one item per filed Pillar video of the sheets folder's
    project whose session has no hand scores."""
    import csv
    import pandas as pd
    from .sheet_sync import SHEET_FOLDER_PROJECT
    if snapshot_dir is None:
        from .config import require
        snapshot_dir = require("snapshot_dir")
    if hand_csv is None:
        from .hand_scores_file import FILE_NAME, out_dir
        hand_csv = out_dir() / FILE_NAME
    with open(hand_csv, newline="", encoding="utf-8") as f:
        scored = {r["video_key"] for r in csv.DictReader(f)}
    v = pd.read_parquet(Path(snapshot_dir) / "video_pellets.parquet",
                        columns=["video_name", "subject_id", "session_date", "tray_type",
                                 "run_number"])
    v = v.drop_duplicates("video_name")
    v = v[v.subject_id.str.startswith(SHEET_FOLDER_PROJECT + "_") & (v.tray_type == "P")]
    out = []
    for r in v.sort_values("video_name").itertuples():
        if r.video_name in scored:
            continue
        out.append({"cohort_id": "_".join(r.subject_id.split("_")[:2]), "source": "database",
                    "says": ("%s is in the database but its session has no hand scores in the "
                             "sheet (animal %s, %s, tray P%d). Check the sheet has that animal, "
                             "date and tray -- and that the video is named for the right ones."
                             % (r.video_name, r.subject_id, str(r.session_date)[:10],
                                int(r.run_number)))})
    return out


def items() -> List[dict]:
    """Everything that needs a person, from every source. A source that
    cannot be read is itself an item, never silence."""
    out: List[dict] = []
    for name, fn in (("reaching pipeline", mousereach_items), ("sheet import", sheet_items),
                     ("database", filed_without_scores)):
        try:
            out += fn()
        except Exception as e:
            out.append({"cohort_id": None, "source": name,
                        "says": "Could not check the %s for things to fix (%s)." % (name, e)})
    return out


def write_fix_these(cohort_id: str, cohort_items: List[dict], sheets_dir: Path) -> Optional[Path]:
    """Write (or, with nothing listed, remove) the cohort's FIX THESE file in
    the sheets folder. Returns the path written, else None. Raises when it
    cannot be written (e.g. open in Excel): the caller reports it."""
    dest = Path(sheets_dir) / (FIX_THESE_NAME % cohort_id)
    if not cohort_items:
        if dest.exists():
            dest.unlink()
        return None
    import shutil
    import tempfile
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fix these"
    ws.append(["%s: %d thing(s) to fix" % (cohort_id, len(cohort_items))])
    ws["A1"].font = Font(bold=True, size=14, color="B71C1C")
    ws.append(["Written by mousedb after the sheet import of %s. Fix each line in the "
               "cohort's tracking sheet (or, where it says so, the video's name). This file "
               "is rewritten after every import and disappears by itself once nothing is "
               "left. Do not edit it -- it is not read." % datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(["What needs fixing", "Found by"])
    for c in "AB":
        ws["%s4" % c].font = Font(bold=True)
    for i in cohort_items:
        ws.append([i.get("says", ""), i.get("source", "")])
    ws.column_dimensions["A"].width = 120
    ws.column_dimensions["B"].width = 20
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    tmp = Path(tempfile.mkdtemp()) / dest.name
    wb.save(tmp)
    shutil.copyfile(tmp, dest)
    shutil.rmtree(tmp.parent, ignore_errors=True)
    return dest


def write_all_fix_these(sheets_dir: Path, cohort_ids: List[str]) -> Dict[str, str]:
    """Write or remove each cohort's FIX THESE file. Returns {cohort: what
    happened} -- 'written (N)', 'none', or 'FAILED: ...'."""
    everything = items()
    out = {}
    for cid in cohort_ids:
        mine = [i for i in everything if i.get("cohort_id") == cid]
        try:
            write_fix_these(cid, mine, sheets_dir)
            out[cid] = "written (%d)" % len(mine) if mine else "none"
        except Exception as e:
            out[cid] = "FAILED: %s: %s" % (type(e).__name__, e)
    return out
