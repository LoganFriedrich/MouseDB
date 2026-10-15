"""Everything a tracking sheet says about each animal, kept column for column.

WHY THIS EXISTS
---------------
A shared ODC-SCI dataset repeats each animal's details (surgery parameters,
drugs and doses, virus, injury level, survival ...) on every row. The sheet
importer only ever kept a handful of those columns in its typed tables (force,
displacement, velocity, dwell, anesthetic), so an export built from the
database could not carry the rest -- even though the lab had typed it all into
the tracking sheets.

This module copies the animal-level tabs of a cohort's workbook into ONE
long table, ``animal_records`` (subject, tab, record number, column name,
value as text), exactly as written. Long, not wide, on purpose: sheets gain and
rename columns over the years, and a long table takes any column without a
schema change. The per-reach ODC export pivots it back out.

What is copied
--------------
* Project cohort workbooks (one row per animal or per surgery, header in row 1):
  ``0a_Metadata``, ``4_Contusion_Injury_Details``, ``5_SC_Injection_Details``.
  A surgery row is copied only when it proves the surgery HAPPENED: the
  animal's body weight on the day is typed in and the row is dated (see
  ``surgery_verdict``). Everything else on a surgery row -- date, type, virus,
  titre, target, depths, drugs, doses, volume formulas, Survived -- can be
  filled in before the day by the blank-sheet generator or by copying the
  previous cohort's workbook, so none of it counts. A row without a weight is a
  planned surgery: not copied. A row that looks done but cannot be taken as
  done (a device reading but no weight, an unreadable weight, a weight but no
  date or a date cell that is not a date) is held: not copied either.
  The surgery importer (importers.ExcelImporter) decides with this same rule
  -- it calls ``surgery_verdict`` -- so the typed surgeries table and these
  records hold the same surgeries; tests/test_animal_records_proof.py fails if
  the two ever disagree on a row.

  Rows left out are NOT announced here. The importer reads the same workbook
  first on every route that runs this module, and it is the one place they
  are announced: a held row every time, a planned row only once it is
  overdue (``PLANNED_SURGERY_GRACE_DAYS``). Saying it here too would put
  every line in front of a person twice.
* Frozen single-letter cohort workbooks: the ``ODC`` tab (ODC-SCI common data
  elements, one row per animal; its header row is found by the "Animal ID"
  cell because rows above it hold notes). Animal numbers are encoded the
  pipeline's way (cohort number = the letter's alphabet position).

An animal cell this module cannot read (not an id of this cohort) is named
ONCE per import, by exactly one of the two readers. Where the importer reads
the same cell -- same tab, same column (``IMPORTER_ANIMAL_ID_HEADINGS``) --
and refuses it, the importer names it and this module says nothing. Anywhere
else (a heading the importer does not read, a cell the importer reads as a
blank or accepts in another spelling, the ODC tab) this module names it,
because nobody else will. A row that is not a performed surgery is never
named for its id: a plan's id is not a problem until the surgery happens.

Each import REPLACES that cohort's rows for the tabs it read, so a correction
made in the sheet replaces the old value instead of piling up beside it. The
sheet import (sheet_sync) runs this only after the importer has SAVED the same
workbook: when the importer saved nothing (it recorded an error and undid the
whole workbook), these rows are left as they were, so the animal details never
describe a workbook the rest of the database does not hold.
Values are text; dates become YYYY-MM-DD; blanks and "?" are not stored.
"""
from __future__ import annotations

import math
import numbers
import re
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from sqlalchemy import Column, DateTime, Integer, String, Text

CNT_TABS = ("0a_Metadata", "4_Contusion_Injury_Details", "5_SC_Injection_Details")
# The scoring tab repeats each animal's sex on every session row (and it is often the
# only place sex is recorded); one Sex record per animal is taken from it.
TRAY_TAB, TRAY_FIELDS = "3b_Manual_Tray", ("Sex",)
ODC_TAB = "ODC"
SUBJECT_COLUMNS = ("Subject_ID", "SubjectID", "Animal", "Animal ID", "Subject", "Mouse ID")

# The headings the sheet IMPORTER (importers.ExcelImporter) reads a row's animal
# id under, on each tab it shares with this module, in the order it tries them.
# WHY this module needs them: an animal cell that cannot be read is named once
# per import. The importer names the cells IT reads ("the animal cell is not an
# animal id", "these animals belong to cohort ..."), so this module stays silent
# about those -- and only those. The importer does not look under every heading
# this module accepts (its 0a_Metadata reader tries 'SubjectID' and 'Animal'
# only), and a cell under any other heading is read by this module alone:
# staying silent there too would leave the row named by nobody. This mirrors
# the importer's readers (_import_metadata, _import_contusion_surgeries,
# _import_surgeries); tests/test_failed_import_leaves_records_alone.py runs the
# real importer under every heading in SUBJECT_COLUMNS and fails if this table
# and the importer part ways.
IMPORTER_ANIMAL_ID_HEADINGS = {
    "0a_Metadata": ("SubjectID", "Animal"),
    "4_Contusion_Injury_Details": ("Subject_ID", "Animal", "SubjectID", "Subject"),
    "5_SC_Injection_Details": ("Subject_ID", "Animal", "SubjectID", "Mouse ID"),
}

# ---------------------------------------------------------------------------
# What proves that a surgery row records a surgery that HAPPENED
# ---------------------------------------------------------------------------
# This is the ONE definition, for this module and for the surgery importer
# (importers.ExcelImporter, which calls surgery_verdict, surgery_date and
# surgery_weight below and keeps no copy of the rule), so the typed surgeries
# table and these records hold the same surgeries. Two copies of the rule were
# kept once and they drifted: inputs like '22 grams', '1e1', a weight under a
# second heading, or a Date column beside a blank Surgery_Date were taken by one
# and refused by the other. tests/test_animal_records_proof.py runs the
# importer's own gate against this one and fails if they ever disagree.
#
# Proof is the animal's body weight typed on the row. WHY the weight and nothing
# else: the animal is weighed on the table to dose its drugs, so that number
# cannot exist before the day. Every other column can. The blank-sheet generator
# and a workbook copied from the previous cohort pre-fill the date, type, virus,
# titre, target, depths, coordinates, drugs and doses; the volume columns are
# formulas that evaluate to 0 on an empty row; and a Survived or a dwell of 0
# survives a partial clearing. Counting any of those as proof is how a whole
# cohort's planned injections were stored as injections that had happened.
#
# The headings the weight may sit under, in order; the first one filled in
# decides. The blank-sheet generator writes the first; the others are how a
# hand-made sheet may head the same column. The importer reads the same list in
# the same order, so a renamed heading cannot make one table take a surgery the
# other refuses.
SURGERY_WEIGHT_COLUMNS = ("Subject_Weight (g)", "Subject_Weight_g", "Subject_Weight",
                          "Weight (g)", "Weight_g")
SURGERY_PROOF_COLUMNS = {
    "4_Contusion_Injury_Details": SURGERY_WEIGHT_COLUMNS,
    "5_SC_Injection_Details": SURGERY_WEIGHT_COLUMNS,
}
# A weighed row must also be dated before it is taken as a surgery: a weight with
# no date is an event whose day is unknown, which no downstream step can place
# (days post-injury, ordering against sessions). The same holds for a date cell
# that does not read as a date ('TBD', or a date pasted with a dash that is not
# a hyphen): copied, it would publish that text as the surgery date, while the
# importer refuses the row because it cannot parse it. Either name is read; the
# first one filled in decides (a blank Surgery_Date column beside a filled Date
# column means the Date), in this module and the importer alike.
SURGERY_DATE_FIELDS = ("Surgery_Date", "Date")
# Contusion device readings. They are NOT proof (a copied workbook carries them
# too), but a row with a reading and no weight is most likely a surgery whose
# weight was never typed. It is held and named for a person to finish, instead
# of being filed quietly with the plans.
SURGERY_READOUT_FIELDS = {"4_Contusion_Injury_Details": ("Actual_kd", "Actual_displacement")}
# The importer names surgeries by kind; this is the tab each kind is read from,
# so it can call surgery_was_performed(row, SURGERY_TABS[kind]).
SURGERY_TABS = {"contusion": "4_Contusion_Injury_Details", "tracing": "5_SC_Injection_Details"}

PERFORMED, PLANNED, HELD = "performed", "planned", "held"

# How many days a DATED planned surgery row may sit without its weight before a
# person is told about it. WHY: a planned row is not a problem -- the sheet is
# filled in ahead of time on purpose, and after the day it is filled in late by
# nature (the weight goes in when someone gets to it, sometimes from a paper
# notebook days later). So a plan is silent while its day is ahead or only just
# past. More than this many days past, it is overdue: either the surgery was
# done and not written down, or it was not done and the date was never updated,
# and either way a person has to look. An UNDATED planned row is never
# announced: it names no day it is late for. The same window as the expectation
# catalogue's after_grace level; this is the one place the number lives.
PLANNED_SURGERY_GRACE_DAYS = 14

# A weight as people type it: '22', '22.4', '22,4' (decimal comma), '22 g'. Only a
# gram unit is accepted -- '22 kg' or '22 mg' is a mistake to show a person, not
# a weight to trust.
_QUANTITY = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(?:g|grams?)?$", re.IGNORECASE)

_ID = re.compile(r"^[A-Z]+_\d{2}_\d{2}$")


def _table():
    """The ORM class, created lazily so importing this module never touches schema order."""
    from .schema import AnimalRecord
    return AnimalRecord


def text_value(v) -> Optional[str]:
    """A cell as the text stored: dates as YYYY-MM-DD, whole floats without '.0',
    None for blank, NaN or '?' (the frozen workbooks use '?' for 'not recorded')."""
    if v is None:
        return None
    if isinstance(v, datetime):
        if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0):
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if v != v:
            return None
        return str(int(v)) if v.is_integer() else repr(v)
    t = str(v).strip()
    return t if t and t != "?" else None


def _subject(value, cohort_id: str) -> Optional[str]:
    t = text_value(value)
    if t is None:
        return None
    t = t.upper().replace("-", "_")
    if re.fullmatch(r"\d{1,2}", t):
        return "%s_%02d" % (cohort_id, int(t))
    return t if _ID.match(t) and t.startswith(cohort_id + "_") else None


def _other_cohorts_id(value, cohort_id: str) -> bool:
    """True when the cell is a well-formed animal id of ANOTHER cohort (an
    animal borrowed between studies, or a row carried over in a copied
    workbook). This module and the importer both refuse such a row, so the
    two never disagree about it. It is named once: by the importer where it
    reads that cell, by this module everywhere else (_records_from_rows)."""
    t = text_value(value)
    if t is None:
        return False
    t = t.upper().replace("-", "_")
    return bool(_ID.match(t)) and not t.startswith(cohort_id + "_")


def _importer_id_column(raw_header, tab: str) -> Optional[int]:
    """The column the sheet importer reads a row's animal id from on ``tab``,
    or None when it reads none of this tab's columns.

    The importer reads ``row.get(first) or row.get(second) ...`` from pandas.
    pandas names each column by its header cell exactly as typed (spaces kept,
    unlike this module's stripped headings), a second column of the same name
    gets a suffix, and a blank cell is NaN -- which counts as filled for
    ``or``. So the importer reads the first of its headings (see
    IMPORTER_ANIMAL_ID_HEADINGS) that is a column at all, in the first column
    of that name."""
    headings = IMPORTER_ANIMAL_ID_HEADINGS.get(tab, ())
    return next((raw_header.index(h) for h in headings if h in raw_header), None)


def _importer_names_cell(value, cohort_id: str, tab: str) -> bool:
    """True when the sheet importer, reading this same animal cell, refuses
    the row and records why -- so the person is already told, by the importer.

    WHY asked of the importer's own id rule (ExcelImporter._resolve_subject_id,
    on an importer built without a database: only its rule is used) instead of
    worked out again here: the two readers parse ids differently (the
    importer takes 'CNT0501' and '005'; this module takes 'CNT-05-01'), and a
    second copy of the importer's rule would drift from it -- repeating the
    line, or worse, silencing it.

    The cell is shown to that rule as the importer receives it: through
    pandas' own cell parser (the one its read_excel uses), which turns text
    such as 'n/a', 'NA' or 'None' into a blank; an Excel error value ('#REF!'
    from a broken formula) reaches the importer as a blank too. A blank is
    named by the importer only on a surgery tab, and only on a weighed row --
    "the weight on the day is typed but the animal id is blank", with the
    sheet row -- which every row this module asks about is (it asks only about
    performed surgeries). On 0a_Metadata the importer skips it without a word.

    Any doubt answers False, so this module names the row: a line said twice
    is a nuisance, a line said by nobody is a row lost without a trace."""
    if isinstance(value, bool):
        # A tick box (TRUE/FALSE) in an animal cell. pandas reads a whole
        # column at once, and among numbers it turns TRUE into 1 -- the
        # importer then files the row under animal 1 without a word -- while
        # judged alone, as below, it stays TRUE and looks refused. Which one
        # the importer saw depends on the cells around it, which this
        # function is not shown, so it is named here.
        return False
    try:
        from openpyxl.cell.cell import ERROR_CODES
        if isinstance(value, str) and value in ERROR_CODES:
            return tab in SURGERY_PROOF_COLUMNS
        from pandas.io.parsers import TextParser
        from .importers import ExcelImporter   # here: importers imports this module
        seen = TextParser([["cell"], [value]], header=0).read().iloc[0, 0]
        if _blank(seen):
            return tab in SURGERY_PROOF_COLUMNS
        if not seen:
            # A 0 or False: the importer's 'a or b' passes over it and reads the
            # NEXT heading's cell instead, so it never judged this one.
            return False
        probe = ExcelImporter.__new__(ExcelImporter)
        return probe._resolve_subject_id(seen, cohort_id, tab) is None
    except Exception:
        return False


def _ascii(text: str) -> str:
    """A warning as plain ASCII. Warnings reach Windows consoles, which crash on
    characters outside their code page, and a warning quotes the sheet: the
    file name (any tail is allowed after the cohort part) and a date cell typed
    as text (an en dash pasted in) can hold any character. Escaped, not dropped,
    so the person still sees exactly what to look for."""
    return text.encode("ascii", "backslashreplace").decode("ascii")


def _blank(v) -> bool:
    """True for an empty cell however it arrives: None, '' or '?' (this module's
    reader), or NaN / NaT / NA (the importer reads rows with pandas)."""
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() in ("", "?")
    try:
        # NaN and NaT are the only values not equal to themselves.
        return bool(v != v)
    except TypeError:
        # pandas.NA refuses to become a bool; it is a blank.
        return True


def _quantity(v) -> Tuple[str, Optional[float]]:
    """('blank' | 'number' | 'unreadable', value) for a cell that should hold a
    positive number, as typed (importer) or as stored here (text)."""
    if _blank(v):
        return "blank", None
    # A tick box or a date typed into a number column is a mistake, not a number
    # (bool is a kind of int in Python, so it must be refused by name).
    if isinstance(v, (bool, datetime, date)):
        return "unreadable", None
    if isinstance(v, numbers.Real):
        x = float(v)
        if math.isinf(x):
            # Nobody weighs an infinite mouse: an 'inf' typed in a weight cell
            # that a reader turned into a number. (NaN was a blank above.)
            return "unreadable", None
    else:
        m = _QUANTITY.match(str(v).strip())
        if not m:
            return "unreadable", None
        x = float(m.group(1).replace(",", "."))
    if x == 0:
        # 0 is what a template default, a cleared row or a formula over an empty
        # row leaves behind -- nothing was measured.
        return "blank", None
    return ("number", x) if x > 0 else ("unreadable", None)


# Words pandas.to_datetime turns into a date that moves with the clock. Typed into
# a date cell they would give the surgery whatever day the import happened to
# run -- an invented date, different on every run -- so they are not dates.
_RELATIVE_DAY_WORDS = {"now", "today", "yesterday", "tomorrow"}


def _text_names_a_whole_day(text: str) -> bool:
    """True when date text spells out its year, month AND day.

    WHY: pandas.to_datetime fills in whatever the text leaves out. A surgery
    time typed into the date cell ('13:00', '10:30 AM') becomes TODAY at that
    time -- the day the import ran, a different day on every run, like the
    words above -- 'Aug 2025' becomes the 1st and '2025' January 1st. Each
    would be stored as a surgery day nobody recorded. Read twice with two
    different fill-in days, such text comes out as two different days; text
    that names its day comes out as the same day both times. (dateutil is the
    parser pandas itself falls back on for text like this.)"""
    from dateutil import parser
    try:
        first = parser.parse(text, default=datetime(2001, 1, 1))
        second = parser.parse(text, default=datetime(2002, 2, 2))
    except (ValueError, OverflowError, TypeError):
        return False
    return first.date() == second.date()


def _surgery_date(row) -> Tuple[str, Optional[str], Optional[date]]:
    """('blank' | 'date' | 'unreadable', the field read, the day) for a surgery
    row's date.

    The first date heading with something in it decides, like the weight. A date
    cell arrives as a date (the importer's pandas row) or as YYYY-MM-DD text (this
    module's reader). Other text is read with pandas.to_datetime, and only when
    it names a whole day (see _text_names_a_whole_day). A bare number is not a
    date: pandas would read it as nanoseconds after 1970, which is not a day
    anybody operated on. The importer stores the day returned here, so the
    day this module copies and the day the surgeries table holds are one day."""
    field = next((f for f in SURGERY_DATE_FIELDS if not _blank(row.get(f))), None)
    if field is None:
        return "blank", None, None
    v = row.get(field)
    if isinstance(v, datetime):  # pandas.Timestamp is a datetime too
        return "date", field, v.date()
    if isinstance(v, date):
        return "date", field, v
    if (isinstance(v, str) and v.strip().lower() not in _RELATIVE_DAY_WORDS
            and _text_names_a_whole_day(v.strip())):
        import pandas as pd  # here, not at the top: only a text date needs it
        try:
            parsed = pd.to_datetime(v.strip())
        except (ValueError, TypeError, OverflowError):
            parsed = None
        # NaT: the text 'NaT' parses to "no date", which is not a day either.
        if parsed is not None and not _blank(parsed):
            return "date", field, parsed.date()
    return "unreadable", field, None


def _weight_field(row, tab: str) -> str:
    """The heading whose cell decides the weight: the first of the tab's weight
    headings with something in it, else the generator's heading (so a reason
    names the column a person should fill)."""
    columns = SURGERY_PROOF_COLUMNS[tab]
    return next((c for c in columns if not _blank(row.get(c))), columns[0])


def surgery_date(row) -> Optional[date]:
    """The day a surgery row is dated, read exactly as surgery_verdict reads it,
    or None when the row has no date or its date cell is not a date.

    The importer stores this day, and the planned-row grace window is counted
    from it, so neither can read a different day from the one the rule judged."""
    return _surgery_date(row)[2]


def surgery_weight(row, tab: str) -> Optional[float]:
    """The body weight on the day in grams, read exactly as surgery_verdict reads
    it, or None when there is none (blank, 0, or not a weight in grams).

    The importer stores this number, so the weight that proved a surgery is the
    weight recorded for it -- '22.4 g' or a weight under a second heading proves
    the surgery and is stored, instead of proving it and being stored as blank."""
    state, grams = _quantity(row.get(_weight_field(row, tab)))
    return grams if state == "number" else None


def planned_surgery_overdue(row, tab: str, today: date) -> Optional[date]:
    """The planned day of a PLANNED surgery row whose day is more than
    PLANNED_SURGERY_GRACE_DAYS before ``today``; None for every other row.

    None therefore means "nothing to tell a person about this plan": the row is
    not a plan, or it is undated (it names no day to be late for), or its day is
    ahead or inside the grace window (the weight is late by nature). See
    PLANNED_SURGERY_GRACE_DAYS for why."""
    if surgery_verdict(row, tab)[0] != PLANNED:
        return None
    day = surgery_date(row)
    if day is None or (today - day).days <= PLANNED_SURGERY_GRACE_DAYS:
        return None
    return day


def surgery_verdict(row, tab: str) -> Tuple[str, str]:
    """Whether one row of a surgery tab records a surgery that happened.

    ``row`` is anything with ``.get(column)``. Both readers pass {heading: cell}
    with each cell as it was typed (a date, a number or text; this module reads
    the workbook directly, the importer reads it with nothing converted -- see
    ExcelImporter._surgery_rows_as_typed); a pandas row or {column: text} is
    judged the same way. Returns (verdict, reason):

    * (PERFORMED, "") -- the body weight on the day is a positive number and the
      row's date reads as a date: the surgery happened.
    * (PLANNED, why)  -- no weight: a planned surgery (or a cleared row). Nothing
      happened yet; the row counts once its weight is typed in.
    * (HELD, why)     -- the row looks like a surgery that happened but cannot be
      taken as one until a person fixes it (reason says what to fix).

    The reasons are plain ASCII phrases for a person, ready to follow "-- ".
    A tab with no proof rule raises KeyError, so a new surgery tab can never be
    let through ungated by accident."""
    # The first heading with something in it decides; with none filled, the
    # reasons name the generator's heading (see _weight_field).
    field = _weight_field(row, tab)
    state, grams = _quantity(row.get(field))
    if state == "unreadable":
        # ascii(): the cell is shown exactly, and a non-ASCII character in it
        # cannot crash a Windows console that prints this warning.
        return HELD, ("its '%s' cell reads %s, which is not a weight in grams; correct the cell"
                      % (field, ascii(text_value(row.get(field)))))
    if state == "blank":
        readings = [f for f in SURGERY_READOUT_FIELDS.get(tab, ())
                    if _quantity(row.get(f))[0] == "number"]
        if readings:
            return HELD, ("it has a device reading (%s) but no body weight in '%s'; the surgery "
                          "looks done but is not taken as done until its weight is typed in"
                          % (", ".join(readings), field))
        return PLANNED, ("no body weight in '%s', so it is a planned surgery, not one that "
                         "happened" % field)
    when, date_field, _ = _surgery_date(row)
    if when == "blank":
        return HELD, ("it is weighed (%g g) but has no Surgery_Date; it is not taken as done "
                      "until its date is typed in" % grams)
    if when == "unreadable":
        return HELD, ("it is weighed (%g g) but its '%s' cell reads %s, which is not a date; "
                      "it is not taken as done until a real date is typed in"
                      % (grams, date_field, ascii(text_value(row.get(date_field)))))
    return PERFORMED, ""


def surgery_was_performed(row, tab: str) -> bool:
    """True only for a row that records a surgery that happened (see surgery_verdict)."""
    return surgery_verdict(row, tab)[0] == PERFORMED


def _records_from_rows(rows, header, tab, cohort_id, subject_col, source, warnings,
                       gated=False, importer_reads_tab=False, raw_header=None) -> List[dict]:
    """Records from one tab's data rows. ``gated``: this is a surgery tab, and a
    row is copied only when surgery_verdict() says the surgery happened.
    ``importer_reads_tab``: the sheet importer reads this same workbook before
    this module and names the rows it refuses (true of the project workbooks'
    tabs, not of the frozen workbooks' ODC tab, which only this module reads).
    ``raw_header``: the header cells as typed, unstripped -- what the importer
    names its columns by (``header`` when not given).

    A surgery row left out is not announced here: the importer judges the same
    rows with the same rule and announces them, on every route that runs this
    module (see the module docstring). Saying it here as well put each planned
    and held row in front of a person twice per import. An animal cell this
    module cannot read is named here unless the importer reads that same cell
    and names it itself (see IMPORTER_ANIMAL_ID_HEADINGS)."""
    out, seen = [], {}
    if gated and not any(c in header for c in SURGERY_PROOF_COLUMNS[tab]):
        # Without a weight column no row on this tab can be shown to have
        # happened, so nothing is copied: copying them all would let plans
        # through. The importer tells the person to restore the heading.
        return out
    # Is the animal column this module reads the very column the importer
    # reads? Only then can the importer have named a cell of it.
    same_column_as_importer = importer_reads_tab and _importer_id_column(
        list(header if raw_header is None else raw_header), tab) == subject_col
    for r in rows:
        if r is None or subject_col >= len(r):
            continue
        # Judged on the cells as the workbook holds them (a date, a number,
        # text), not on the text copied below: the importer judges its own read
        # of the same cells as typed, and a cell turned into text first can read
        # differently (the number 20250827 is not a date; the text '20250827'
        # reads as one), which would let the two tables disagree. A heading that
        # appears twice is judged by its FIRST column, as the importer's reader
        # (pandas) names the second one 'X.1'.
        raw = {}
        for i, v in enumerate(r):
            if i < len(header) and header[i] and i != subject_col:
                raw.setdefault(header[i], v)
        performed = (not gated) or surgery_verdict(raw, tab)[0] == PERFORMED
        sid = _subject(r[subject_col], cohort_id)
        if sid is None:
            v = r[subject_col]
            if text_value(v) is None:
                continue
            if not performed:
                # A plan, or a row the importer names as held: either way not
                # a row this module could have copied, and not one to announce
                # for its id (an id seen only on a plan is not a problem).
                continue
            if same_column_as_importer and _importer_names_cell(v, cohort_id, tab):
                # The importer read this very cell, refused the row and named
                # it ("not an animal id", or "belong to cohort ..."); a line
                # here said the same thing a second time.
                continue
            # Named here: under a heading the importer does not read, or a cell
            # the importer read as a blank or as another, valid, id -- either
            # way no other line tells the person about this row.
            # ascii(), not repr(): repr keeps a non-ASCII character as it is,
            # and printing that on a Windows console crashes the import.
            if _other_cohorts_id(v, cohort_id):
                home = text_value(v).upper().replace("-", "_").rsplit("_", 1)[0]
                warnings.append(_ascii(
                    "%s: '%s' row for animal %s skipped: it belongs to cohort %s, not to "
                    "this workbook's cohort %s, and an animal's details are copied only "
                    "from its own cohort's workbook" % (source, tab, ascii(v), home, cohort_id)))
            else:
                warnings.append(_ascii("%s: '%s' row with unrecognised animal %s skipped"
                                       % (source, tab, ascii(v))))
            continue
        cells = {header[i]: text_value(v) for i, v in enumerate(r)
                 if i < len(header) and header[i] and i != subject_col}
        cells = {k: v for k, v in cells.items() if v is not None}
        # A row holding nothing but the animal id has nothing to copy, so
        # leaving it out drops nothing and is not worth a warning.
        if not cells or not performed:
            continue
        n = seen[sid] = seen.get(sid, 0) + 1
        for field, value in cells.items():
            out.append({"subject_id": sid, "cohort_id": cohort_id, "source_tab": tab,
                        "record_no": n, "field": field, "value": value, "source_file": source})
    return out


def read_cnt_workbook(path: Path, cohort_id: str, warnings: List[str]) -> List[dict]:
    """Records from a project cohort workbook's animal-level tabs (header in row 1)."""
    import openpyxl
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: List[dict] = []
    try:
        for tab in CNT_TABS:
            if tab not in wb.sheetnames:
                continue
            it = wb[tab].iter_rows(values_only=True)
            try:
                raw_header = list(next(it))
            except StopIteration:
                continue
            header = [text_value(h) or "" for h in raw_header]
            col = next((header.index(c) for c in SUBJECT_COLUMNS if c in header), None)
            if col is None:
                warnings.append(_ascii("%s: '%s' has no animal id column" % (path.name, tab)))
                continue
            out += _records_from_rows(it, header, tab, cohort_id, col, path.name, warnings,
                                      gated=tab in SURGERY_PROOF_COLUMNS,
                                      importer_reads_tab=True, raw_header=raw_header)
        if TRAY_TAB in wb.sheetnames:
            out += _tray_fields(wb[TRAY_TAB], cohort_id, path.name)
    finally:
        wb.close()
    return out


def _tray_fields(ws, cohort_id: str, source: str) -> List[dict]:
    """First non-blank value of each TRAY_FIELDS column per animal on the scoring tab."""
    it = ws.iter_rows(values_only=True)
    try:
        header = [text_value(h) or "" for h in next(it)]
    except StopIteration:
        return []
    col = next((header.index(c) for c in SUBJECT_COLUMNS if c in header), None)
    fields = [(header.index(f), f) for f in TRAY_FIELDS if f in header]
    if col is None or not fields:
        return []
    found: Dict[tuple, str] = {}
    for r in it:
        if r is None or col >= len(r):
            continue
        sid = _subject(r[col], cohort_id)
        if sid is None:
            continue
        for i, f in fields:
            v = text_value(r[i]) if i < len(r) else None
            if v is not None and (sid, f) not in found:
                found[(sid, f)] = v
    return [{"subject_id": sid, "cohort_id": cohort_id, "source_tab": TRAY_TAB, "record_no": 1,
             "field": f, "value": v, "source_file": source} for (sid, f), v in found.items()]


def read_odc_tab(path: Path, cohort_id: str, warnings: List[str]) -> List[dict]:
    """Records from a frozen letter-cohort workbook's ODC tab (header row found by 'Animal ID')."""
    import openpyxl
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if ODC_TAB not in wb.sheetnames:
            return []
        rows = list(wb[ODC_TAB].iter_rows(values_only=True))
    finally:
        wb.close()
    for i, r in enumerate(rows[:15]):
        header = [text_value(h) or "" for h in (r or ())]
        if "Animal ID" in header:
            col = header.index("Animal ID")
            return _records_from_rows(rows[i + 1:], header, ODC_TAB, cohort_id, col, path.name, warnings)
    warnings.append(_ascii("%s: '%s' tab has no 'Animal ID' header" % (path.name, ODC_TAB)))
    return []


def ensure_table(db) -> None:
    """Create animal_records if this database predates it (adds a table; changes nothing else)."""
    _table().__table__.create(db.engine, checkfirst=True)


def replace_cohort(db, cohort_id: str, tabs, records: List[dict]) -> int:
    """Swap this cohort's rows for ``tabs`` with ``records`` in one transaction."""
    AnimalRecord = _table()
    ensure_table(db)
    now = datetime.now()
    with db.session() as s:
        s.query(AnimalRecord).filter(AnimalRecord.cohort_id == cohort_id,
                                     AnimalRecord.source_tab.in_(list(tabs))
                                     ).delete(synchronize_session=False)
        s.bulk_insert_mappings(AnimalRecord, [dict(r, imported_at=now) for r in records])
    return len(records)


def sync_cohort(db, cohort_id: str, sheet: Path, kind: str = "cohort", dry_run: bool = False) -> dict:
    """Read and (unless dry_run) store one cohort's animal records. Never raises:
    returns {"records": n, "warnings": [...], "error": text-or-None}."""
    warnings: List[str] = []
    try:
        if kind == "odc":
            records, tabs = read_odc_tab(sheet, cohort_id, warnings), (ODC_TAB,)
        else:
            records, tabs = read_cnt_workbook(sheet, cohort_id, warnings), CNT_TABS + (TRAY_TAB,)
        n = len(records) if dry_run else replace_cohort(db, cohort_id, tabs, records)
        return {"records": n, "warnings": warnings, "error": None}
    except Exception as e:
        return {"records": 0, "warnings": warnings, "error": "%s: %s" % (type(e).__name__, e)}
